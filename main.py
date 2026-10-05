"""
BSF Saudi Legal Case Workbench — Standard WebApp Backend
File: backend.py
"""

import io
import re
import json
import math
import mimetypes
import uuid
import base64
import threading
import time
import traceback
import os
import sys

import pandas as pd
from flask import request, jsonify, send_file, Response, copy_current_request_context, g
from dataiku.customwebapp import *  
import dataiku

# -----------------------------------------------------------------------------
# USAGE ANALYTICS & MONITORING
# -----------------------------------------------------------------------------
sys.path.insert(0, "/dataiku/design/plugins/dev/usage-analytics-lib/python-lib")
try:
    from usageanalyticslib import write_usage_event
except ImportError:
    def write_usage_event(*args, **kwargs):
        pass

PROJECT_KEY = os.environ.get("DKU_CURRENT_PROJECT_KEY", "unknown_project")
LITIGATION_PROJECT_KEY = f"{PROJECT_KEY}_LITIGATION"
AGREEMENT_PROJECT_KEY = f"{PROJECT_KEY}_AGREEMENT"

def increment_usage(session_id, field, amount=1, use_case="litigation"):
    if not session_id:
        return
    virtual_key = AGREEMENT_PROJECT_KEY if use_case == "agreement" else LITIGATION_PROJECT_KEY
    try:
        write_usage_event(virtual_key, session_id, field, amount)
    except Exception:
        pass


# -----------------------------------------------------------------------------
# CORE LEGAL PLATFORM IMPORTS
# -----------------------------------------------------------------------------
from legal_platform.intake import create_case, add_message
from legal_platform import access
from legal_platform.storage import list_cases, latest_case, case_rows, replace_case_rows
from legal_platform.workflow import (
    COMPLETED,
    ERROR,
    NOT_STARTED,
    READY,
    RUNNING,
    WAITING,
    build_stage_overview,
    read_workflow_state,
    stage_patch,
    update_workflow_state,
)
from legal_platform.files import save_upload
from legal_platform.extraction import extract_pdf_page_by_page, load_page_extraction
from legal_platform.case_mapping import build_case_map
from legal_platform.case_map_storage import persist_case_map
from legal_platform.case_register import build_case_register, prioritise_pages_for_summary
from legal_platform.facts import add_fact_candidate
from legal_platform.chat_orchestrator import assess_next_step, run_legal_analysis
from legal_platform.interview_state import persist_interview_state
from legal_platform.case_summary import generate_case_summary, approve_case_summary
from legal_platform.attorney_workbench import answer_case_question, revise_bilingual_pleading
from legal_platform.audit import audit
from legal_platform.agreement_workbench import (
    AGREEMENT_TYPES, RELATIONSHIP_TYPES, classify_agreement, extract_clause_map,
    run_agreement_review, save_agreement_state, load_agreement_states,
    discuss_agreement,
)

# -----------------------------------------------------------------------------
# ACCOUNTING & FORENSIC DISPUTE ENGINE IMPORTS
# -----------------------------------------------------------------------------
from legal_platform.config import (
    FINANCIAL_DISCREPANCIES_DATASET,
    FINANCIAL_DOCUMENT_CLASSIFICATION_DATASET,
    FINANCIAL_FINDINGS_DATASET,
    FINANCIAL_LINE_ITEM_CORRECTIONS_DATASET,
    FINANCIAL_LINE_ITEMS_DATASET,
    FINANCIAL_TIMELINE_DATASET,
    CASE_DOCUMENT_FOLDER_ID,
)
from legal_platform.financial_classification import classify_case_pages, pages_needing_classification
from legal_platform.financial_corrections import (
    load_fact_resolutions,
    load_latest_corrections,
    submit_correction,
    submit_fact_resolution,
    submit_fact_resolutions,
)
from legal_platform.financial_facts import (
    PROPOSAL_ACTION,
    format_amount,
    parse_amount,
    UNCERTAIN,
    build_fact_ledger,
    build_fact_review_items,
    build_fact_table,
    fact_from_row,
    user_fact_row,
    count_uncertain,
    effective_fact,
    validate_fact_fields,
)
from legal_platform.financial_reconciliation import persist_reconciled_rows, reinterpret_fact
from legal_platform.financial_fields import FIELD_NAMES
from legal_platform.financial_extraction_pipeline import (
    load_document_claims,
    run_financial_extraction,
    save_document_claims,
)
from legal_platform.financial_forensics import (
    build_and_save_financial_timeline,
    ledger_summary,
    load_saved_forensic_results,
    run_claim_based_accounting_analysis,
)
from legal_platform.pleading import (
    build_drafting_record,
    draft_pleading,
    is_full_pleading,
    pleading_language,
    pleading_languages,
    pleading_to_markdown,
    revise_pleading,
)
from legal_platform.financial_normalizer import (
    normalize_amount,
    normalize_debit_credit,
)

APP_ASSETS_FOLDER_ID = "qx2RWzgX"
APP_LOGO_PATH = "logo.png"

RUN_JOBS_INLINE = False
JOBS = {}
JOBS_LOCK = threading.Lock()

# =============================================================================
# TESTING WORKFLOW BYPASS
# =============================================================================
# False = a prepared attorney summary is enough to continue downstream.
# True  = require formal attorney approval and a clean case.
# IMPORTANT: set this to True before production.
REQUIRE_APPROVED_SUMMARY = False


def approval_gate_passed(state):
    """
    Keep the real workflow dependencies while optionally bypassing formal
    attorney approval during testing.

    Testing mode (REQUIRE_APPROVED_SUMMARY = False):
        - an attorney summary must exist
        - formal approval is not required
        - case_dirty does not block downstream testing

    Production mode (REQUIRE_APPROVED_SUMMARY = True):
        - an attorney summary must exist
        - it must be formally approved
        - the case must not be dirty
    """
    if not state or state.get("attorney_summary") is None:
        return False

    if not REQUIRE_APPROVED_SUMMARY:
        return True

    return (
        bool(state.get("summary_approved"))
        and not bool(state.get("case_dirty"))
    )


# =============================================================================
# SANITIZATION HELPERS
# =============================================================================
def _clean_scalar(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return str(value)
    return value


def _json_safe(value):
    if value is None or isinstance(value, (str, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, int):
        return value
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_json_safe(item) for item in value]
    if isinstance(value, pd.Timestamp):
        return str(value)
    if hasattr(value, "item"):
        try:
            return _json_safe(value.item())
        except (AttributeError, ValueError):
            pass
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return str(value)


def _ok(payload):
    return jsonify(_json_safe(payload))


def frame_to_records(frame):
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        return []
    return json.loads(frame.fillna("").to_json(orient="records", force_ascii=False))


# =============================================================================
# ACCOUNTING ANALYSIS HELPERS
# =============================================================================
def _page_id_column(pages_df):
    return "case_document_page_id" if "case_document_page_id" in pages_df.columns else "page_id"


def _normalized_ledger(case_id, line_items_df, corrections=None, resolutions=None):
    """(ledger, withheld_count). The final normalized ledger: every atomic
    fact whose status is final (extracted, calculated, inferred, missing,
    user-confirmed, user-corrected). UNCERTAIN facts are withheld until the
    user resolves them; IGNORED facts are left out."""
    if line_items_df is None or line_items_df.empty:
        return [], 0
    if corrections is None:
        corrections = load_latest_corrections(case_id)
    if resolutions is None:
        resolutions = load_fact_resolutions(case_id)
    return build_fact_ledger(line_items_df.to_dict(orient="records"), corrections, resolutions)


def _pages_by_id(data):
    pages_by_id = {}
    pages_df = data.get("pages")
    if isinstance(pages_df, pd.DataFrame) and not pages_df.empty:
        labels = _page_reference_map(data)
        id_col = _page_id_column(pages_df)
        for row in pages_df.fillna("").to_dict(orient="records"):
            page_id = str(row.get(id_col, "") or row.get("page_id", ""))
            if page_id:
                pages_by_id[page_id] = {
                    "page_text": str(row.get("page_text", "") or ""),
                    "label": labels.get(page_id, ""),
                }
    return pages_by_id


def _review_items(case_id, data, corrections=None, resolutions=None):
    """UNCERTAIN facts waiting for the user, each with the page evidence
    needed to decide it (see financial_facts.build_fact_review_items)."""
    line_items = data.get("financial_line_items")
    if not isinstance(line_items, pd.DataFrame) or line_items.empty:
        return []
    if corrections is None:
        corrections = load_latest_corrections(case_id)
    if resolutions is None:
        resolutions = load_fact_resolutions(case_id)
    items = build_fact_review_items(
        line_items.to_dict(orient="records"), corrections, resolutions, _pages_by_id(data),
    )
    for item in items:
        item["page_image_url"] = f"/page_image?case_id={case_id}&page_id={item['page_id']}"
    return items


# =============================================================================
# CASE DATA LOADING
# =============================================================================
def load_case_data(case_id):
    data = {
        "case": latest_case(case_id) or {"case_id": case_id},
        "messages": case_rows("case_messages", case_id),
        "documents": case_rows("case_documents", case_id),
        "pages": case_rows("case_document_pages", case_id),
        "facts": case_rows("case_facts", case_id),
        "fact_candidates": case_rows("case_fact_candidates", case_id),
        "parties": case_rows("case_parties", case_id),
        "events": case_rows("case_events", case_id),
        "contradictions": case_rows("case_contradictions", case_id),
        "issues": case_rows("case_issues", case_id),
        "issue_candidates": case_rows("case_issue_candidates", case_id),
        "evidence": case_rows("case_evidence", case_id),
        # Workflow state is read per case from its state file (see
        # legal_platform.workflow); audit_events and case_approvals are not
        # loaded here: they hold every case's history and are large.
        "financial_line_items": case_rows(FINANCIAL_LINE_ITEMS_DATASET, case_id),
        "financial_classifications": case_rows(FINANCIAL_DOCUMENT_CLASSIFICATION_DATASET, case_id),
    }
    data["financial_line_items"] = _facts_on_financial_pages(data)

    forensic = load_saved_forensic_results(case_id)
    data["financial_timeline"] = forensic.get("timeline", [])
    data["forensic_findings"] = {"claim_evaluations": forensic.get("claim_evaluations", [])}
    data["forensic_run_metadata"] = forensic.get("run_metadata", {})
    data["discrepancies"] = forensic.get("discrepancies", [])
    data["cross_check_summary"] = forensic.get("cross_check_summary", {})

    return data


def _facts_on_financial_pages(data):
    """Facts of pages classified financial or mixed (or not classified yet). A page
    later classified as an email, letter or claim keeps its rows on record,
    but they leave the review table and the ledger. Facts the user added
    are always kept."""
    frame = data.get("financial_line_items")
    if not isinstance(frame, pd.DataFrame) or frame.empty or "page_id" not in frame.columns:
        return frame
    latest = _latest_classifications_by_page(data)
    excluded = {pid for pid, info in latest.items()
                if info.get("status") != "failed" and info.get("page_type") and info["page_type"] not in {"financial", "mixed"}}
    if not excluded:
        return frame
    user_added = frame["row_id"].astype(str).str.startswith("FACTU") if "row_id" in frame.columns else False
    return frame[~frame["page_id"].astype(str).isin(excluded) | user_added].copy()


def _latest_classifications_by_page(data):
    frame = data.get("financial_classifications")
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        return {}
    working = frame.sort_values("created_at") if "created_at" in frame.columns else frame
    latest = {}
    for row in working.fillna("").to_dict(orient="records"):
        page_id = str(row.get("page_id", "") or row.get("case_document_page_id", ""))
        if not page_id:
            continue
        latest[page_id] = {
            "page_type": str(row.get("page_type", "") or row.get("document_type", "")).lower().strip(),
            "confidence": _clean_scalar(row.get("confidence")),
            "status": str(row.get("classification_status", "") or "completed"),
        }
    return latest


def best(data, approved_key, candidate_key):
    return data[approved_key] if not data[approved_key].empty else data[candidate_key]


# =============================================================================
# PERSISTENCE & WORKFLOW STATE
#
# State lives in legal_platform.workflow: every update is merged into the
# latest stored state (append-only, one audit row per update) so parallel
# jobs keep each other's results. Routes pass only the keys they change.
# =============================================================================
def restore_workflow_state(data):
    return read_workflow_state(str((data.get("case") or {}).get("case_id", "")))


def invalidation_changes(reason):
    """State changes applied when new material arrives. Generated content
    (review, analysis, pleading...) is kept and shown as stale; formal
    approval is revoked so nothing proceeds on an outdated sign-off."""
    return {
        "case_dirty": True,
        "dirty_reason": reason,
        "accounting_dirty": True,
        "accounting_dirty_reason": reason,
        "summary_approved": False,
        "summary_approved_by": "",
    }


def accounting_state(data, state, review_items=None):
    """
    Current accounting stage from the explicit workflow state plus the
    persisted accounting data.

    An accounting run may legitimately produce zero financial line items,
    so presence of line items must NOT be used as the completion flag —
    accounting_status (set explicitly by the /accounting/* routes below)
    is the single source of truth for "what stage are we at". Pending
    review items always win: while any exist the stage waits for the user.
    """
    status = str(state.get("accounting_status", "not_started") or "not_started").strip().lower()

    line_items = data.get("financial_line_items")
    has_line_items = isinstance(line_items, pd.DataFrame) and not line_items.empty

    classifications = data.get("financial_classifications")
    has_classifications = isinstance(classifications, pd.DataFrame) and not classifications.empty

    if review_items is None:
        review_items = _review_items(str(data["case"].get("case_id", "")), data)
    pending_count = len(review_items)

    findings = data.get("forensic_findings") or {}
    has_forensic_output = bool(findings.get("claim_evaluations") or data.get("discrepancies"))

    return {
        "status": status,
        "dirty": bool(state.get("accounting_dirty")),
        "dirty_reason": state.get("accounting_dirty_reason", ""),
        "has_classifications": has_classifications,
        "has_line_items": has_line_items,
        "has_pending_review": bool(pending_count),
        "pending_count": pending_count,
        "pending_field_count": len(review_items),
        "has_forensic_output": has_forensic_output,
        "facts_confirmed": bool(state.get("accounting_facts_confirmed")),
    }


def stage_overview(case_id, state, data, accounting=None, register=None):
    """What every stage is doing: stored results + stage records + live jobs."""
    if accounting is None:
        accounting = accounting_state(data, state)
    if register is None:
        register = case_register(data)
    return build_stage_overview({
        "has_documents": not data["documents"].empty and not data["pages"].empty,
        "has_summary": state.get("attorney_summary") is not None,
        "gate_passed": approval_gate_passed(state),
        "case_dirty": bool(state.get("case_dirty")),
        "has_analysis": state.get("analysis") is not None,
        "has_memo": state.get("memo") is not None,
        "pleading_status": state.get("pleading_status", "draft"),
        "accounting": accounting,
        "stages": state.get("stages") or {},
        "active_jobs": _active_jobs(case_id),
        "register_fingerprint": register.get("fingerprint"),
    })


def _party_rank(role):
    text = str(role or "").lower()
    priorities = [
        (0, ("claimant", "plaintiff", "applicant", "مدعي", "طالب")),
        (1, ("defendant", "respondent", "مدعى عليه", "مدعى عليها")),
        (2, ("appellant", "مستأنف")),
        (3, ("bank", "بنك", "مصرف")),
    ]
    for rank, words in priorities:
        if any(word in text for word in words):
            return rank
    return 99


def _ordered_parties(items):
    return sorted(
        list(items or []),
        key=lambda item: (_party_rank(item.get("role")), str(item.get("name", "")).lower()),
    )


def _chronology_key(item):
    raw = str(item.get("date", "") or "").strip()
    if not raw or str(item.get("date_precision", "")).lower() == "unknown":
        return (1, "9999-99-99", raw)
    return (0, raw.replace("/", "-"), raw)


def _ordered_chronology(items):
    return sorted(list(items or []), key=_chronology_key)


def _page_reference_map(data):
    documents = {}
    doc_frame = data.get("documents")
    if isinstance(doc_frame, pd.DataFrame) and not doc_frame.empty:
        for _, row in doc_frame.fillna("").iterrows():
            doc_id = str(row.get("case_document_id", "") or row.get("document_id", "") or "")
            filename = str(row.get("original_filename", "") or row.get("file_name", "") or "Document")
            if doc_id:
                documents[doc_id] = filename

    mapping = {}
    page_frame = data.get("pages")
    if isinstance(page_frame, pd.DataFrame) and not page_frame.empty:
        for _, row in page_frame.fillna("").iterrows():
            page_id = str(row.get("case_document_page_id", "") or row.get("page_id", "") or "")
            doc_id = str(row.get("case_document_id", "") or row.get("document_id", "") or "")
            page_number = row.get("page_number", "")
            filename = documents.get(doc_id, str(row.get("original_filename", "") or "Document"))
            label = f"{filename} — page {page_number}" if page_number != "" else filename
            if page_id:
                mapping[page_id] = label
    return mapping


def _page_labels(source_ids, data):
    mapping = _page_reference_map(data)
    labels = []
    for source_id in source_ids or []:
        value = str(source_id or "").strip()
        if not value:
            continue
        label = mapping.get(value, f"Page {value[:6]}")
        if label not in labels:
            labels.append(label)
    return labels


def case_register(data):
    """The unified case register (see legal_platform.case_register): one
    consolidated view of parties, events, facts, issues, evidence requests
    and contradictions across every processed document."""
    return build_case_register(data, _page_reference_map(data))


def short_case_reference(case_id):
    raw = re.sub(r"[^A-Za-z0-9]", "", str(case_id or ""))
    suffix = raw[-6:].upper() if raw else "------"
    return f"CASE-{suffix}"


def first_case_value(row, *columns, default="—"):
    for column in columns:
        value = row.get(column, "")
        if value is not None and str(value).strip() and str(value).lower() != "nan":
            return str(value).strip()
    return default


def case_search_frame(cases, query):
    if cases.empty or not query.strip():
        return cases
    needle = query.strip().casefold()
    searchable = [
        "case_name", "client_name", "customer_name", "matter_type",
        "responsible_attorney", "created_by", "case_id", "case_status",
    ]
    mask = pd.Series(False, index=cases.index)
    for column in searchable:
        if column in cases.columns:
            mask = mask | cases[column].fillna("").astype(str).str.casefold().str.contains(needle, regex=False)
    if "case_id" in cases.columns:
        references = cases["case_id"].fillna("").astype(str).map(short_case_reference).str.casefold()
        mask = mask | references.str.contains(needle, regex=False)
    return cases[mask]


def filter_cases_by_workflow(cases, workflow_type):
    if cases.empty:
        return cases
    frame = cases.copy()
    if "workflow_type" not in frame.columns:
        return frame if workflow_type == "litigation" else frame.iloc[0:0]
    values = frame["workflow_type"].fillna("").astype(str).str.strip().str.lower()
    if workflow_type == "litigation":
        return frame[(values == "") | (values == "litigation")]
    return frame[values == "agreement_review"]


def serialise_case_card(item, workflow_type):
    case_id = str(item.get("case_id", ""))
    activity = first_case_value(item, "updated_at", "created_at", default="—")
    if "T" in activity:
        activity = activity.replace("T", " ")[:16]
    status = first_case_value(item, "workflow_status", "case_status", default="")
    return {
        "case_id": case_id,
        "reference": short_case_reference(case_id),
        "title": first_case_value(item, "case_name", "client_name", "customer_name", default=""),
        "matter": first_case_value(item, "matter_type", default=""),
        "attorney": first_case_value(item, "responsible_attorney", "created_by", default=""),
        "activity": activity,
        "status": status,
        "workflow_type": workflow_type,
    }


def serialise_summary_support(summary, data):
    """Preserves exactly what the LLM generated into standard arrays."""
    if not isinstance(summary, dict):
        return {}

    chronology_list = summary.get("chronology") or []
    parties_list = summary.get("parties") or []

    contradictions_list = []
    contradictions_frame = data.get("contradictions")
    if isinstance(contradictions_frame, pd.DataFrame) and not contradictions_frame.empty:
        for row in contradictions_frame.fillna("").to_dict(orient="records"):
            try:
                source_page_ids = json.loads(row.get("source_page_ids_json", "[]") or "[]")
                if not isinstance(source_page_ids, list):
                    source_page_ids = []
            except (TypeError, json.JSONDecodeError):
                source_page_ids = []
            contradictions_list.append({
                "description": row.get("description", ""),
                "clarification_required": row.get("clarification_required", ""),
                "source_page_ids": source_page_ids,
                "page_labels": _page_labels(source_page_ids, data),
            })

    return {
        "ordered_parties": _ordered_parties(parties_list),
        "ordered_chronology": [
            {"item": item, "page_labels": _page_labels(item.get("source_page_ids", []), data)}
            for item in _ordered_chronology(chronology_list)
        ],
        "contradictions": contradictions_list,
    }

def memo_to_markdown(memo, language):
    if not isinstance(memo, dict):
        return str(memo or "")
    if is_full_pleading(memo):
        return pleading_to_markdown(memo, language)
    section = memo.get("pleading_ar" if language == "ar" else "pleading_en", {})
    if isinstance(section, str):
        return section

    title = memo.get("title_ar" if language == "ar" else "title_en", "")
    filing_type = memo.get("filing_type_ar" if language == "ar" else "filing_type_en", "")
    is_ar = language == "ar"
    lines = []

    if is_ar and section.get("basmala"):
        lines.append(section.get("basmala"))
    if filing_type:
        lines.append(f"# {filing_type}")
    if title:
        lines.append(f"## {title}")

    for key in ("court_heading", "case_details", "party_heading", "subject", "formal_salutation", "opening"):
        value = section.get(key, "")
        if value:
            lines.append(str(value))

    def _fmt(sec):
        if isinstance(sec, str): return sec
        if isinstance(sec, list):
            res = []
            for it in sec:
                if isinstance(it, dict):
                    d = it.get("date", "")
                    f = it.get("fact") or it.get("supported_fact") or it.get("text", "")
                    prefix = f"**{d}:** " if d else "- "
                    res.append(f"{prefix}{f}")
                else:
                    res.append(f"- {str(it)}")
            return "\n".join(res)
        return str(sec or "")

    if section.get("facts"):
        lines.append("## أولاً: الوقائع وتتبع حركة الأموال" if is_ar else "## I. Statement of Facts & Fund Flows")
        lines.append(_fmt(section.get("facts")))
    if section.get("procedural_defences"):
        lines.append("## ثانياً: الدفوع الشكلية والإجرائية" if is_ar else "## II. Procedural Defences")
        lines.append(_fmt(section.get("procedural_defences")))
    if section.get("substantive_defences"):
        lines.append("## ثالثاً: الدفوع الموضوعية والنظامية" if is_ar else "## III. Substantive Defences")
        lines.append(_fmt(section.get("substantive_defences")))
    if section.get("response_to_opponent"):
        lines.append("## رابعاً: الرد على ادعاءات الخصم" if is_ar else "## IV. Response to Opposing Party")
        lines.append(_fmt(section.get("response_to_opponent")))
    if section.get("requests"):
        lines.append("## خامساً: الطلبات الختامية" if is_ar else "## V. Relief Requested")
        lines.append(_fmt(section.get("requests")))
    if section.get("closing"):
        lines.append(section.get("closing"))
    if section.get("signature_block"):
        lines.append(section.get("signature_block"))

    return "\n\n".join(lines)


def build_pleading_docx_bytes(memo, case_reference="", languages=("ar", "en")):
    """Word version of the pleading, in the given languages (Arabic
    right-to-left). Headings, tables, lists and **bold** labels of the
    markdown are kept."""
    from io import BytesIO
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement

    def _set_rtl(paragraph):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_pr = paragraph._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)

    def _add_runs(paragraph, text):
        for index, part in enumerate(re.split(r"\*\*(.+?)\*\*", text.replace("`", ""))):
            if part:
                paragraph.add_run(part).bold = index % 2 == 1

    def _add_table(document, rows, rtl):
        cells = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
        cells = [row for row in cells if not all(re.fullmatch(r"-{3,}", cell) for cell in row)]
        if not cells:
            return
        table = document.add_table(rows=len(cells), cols=len(cells[0]))
        table.style = "Table Grid"
        for r, row in enumerate(cells):
            for c, value in enumerate(row[:len(cells[0])]):
                paragraph = table.cell(r, c).paragraphs[0]
                _add_runs(paragraph, f"**{value}**" if r == 0 else value)
                if rtl:
                    _set_rtl(paragraph)
        if rtl:
            tbl_pr = table._tbl.tblPr
            tbl_pr.append(OxmlElement("w:bidiVisual"))

    def _add_markdown_text(document, markdown_text, rtl=False):
        table_rows = []
        for raw_line in str(markdown_text or "").split("\n") + [""]:
            line = raw_line.strip()
            if line.startswith("|"):
                table_rows.append(line)
                continue
            if table_rows:
                _add_table(document, table_rows, rtl)
                table_rows = []
            if not line or line == "---":
                continue
            if line.startswith("### "):
                paragraph = document.add_heading(line[4:], level=3)
            elif line.startswith("## "):
                paragraph = document.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                paragraph = document.add_heading(line[2:], level=1)
            elif line.startswith("- "):
                paragraph = document.add_paragraph(style="List Bullet")
                _add_runs(paragraph, line[2:])
            else:
                paragraph = document.add_paragraph()
                _add_runs(paragraph, line)
            if rtl:
                _set_rtl(paragraph)

    document = Document()
    if case_reference:
        title_paragraph = document.add_paragraph(case_reference)
        title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for index, language in enumerate(languages):
        if index:
            document.add_page_break()
        _add_markdown_text(document, memo_to_markdown(memo, language), rtl=language == "ar")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _load_app_logo_data_uri():
    try:
        folder = dataiku.Folder(APP_ASSETS_FOLDER_ID)
        with folder.get_download_stream(APP_LOGO_PATH) as stream:
            raw = stream.read()
        return "data:image/png;base64," + base64.b64encode(raw).decode("ascii")
    except Exception:
        return ""


def _load_page_image_bytes(page_image_path):
    if not page_image_path:
        return None
    try:
        folder = dataiku.Folder(CASE_DOCUMENT_FOLDER_ID)
        with folder.get_download_stream(page_image_path) as stream:
            return stream.read()
    except Exception:
        return None


def _page_row(data, page_id):
    pages = data.get("pages")
    if not isinstance(pages, pd.DataFrame) or pages.empty:
        return None
    id_column = "case_document_page_id" if "case_document_page_id" in pages.columns else "page_id"
    if id_column not in pages.columns:
        return None
    match = pages[pages[id_column].astype(str) == str(page_id)]
    if match.empty:
        return None
    return match.iloc[0].to_dict()


def serialise_facts_register(data):
    facts_register = best(data, "facts", "fact_candidates")
    rows = []
    if not isinstance(facts_register, pd.DataFrame) or facts_register.empty:
        return rows
    pages_frame = data.get("pages")
    for _, row in facts_register.fillna("").iterrows():
        fact_text = str(row.get("fact_text", ""))
        if not fact_text:
            continue
        fact_id = str(row.get("fact_id", "") or row.get("fact_candidate_id", ""))

        try:
            page_ids = json.loads(row.get("source_page_ids_json", "[]") or "[]")
            if not isinstance(page_ids, list):
                page_ids = []
        except (TypeError, json.JSONDecodeError):
            page_ids = []

        if not page_ids and str(row.get("page_number", "")).strip() and isinstance(pages_frame, pd.DataFrame) and not pages_frame.empty:
            doc_col = "case_document_id" if "case_document_id" in pages_frame.columns else None
            if doc_col and str(row.get("case_document_id", "")).strip():
                match = pages_frame[
                    (pages_frame[doc_col].astype(str) == str(row.get("case_document_id", "")))
                    & (pages_frame["page_number"].astype(str) == str(row.get("page_number", "")))
                ]
                id_col = "case_document_page_id" if "case_document_page_id" in pages_frame.columns else "page_id"
                if not match.empty and id_col in match.columns:
                    page_ids = [str(match.iloc[0][id_col])]

        rows.append({
            "fact_id": fact_id,
            "fact_text": fact_text,
            "confidence": _clean_scalar(row.get("confidence")),
            "source_type": str(row.get("source_type", "") or ""),
            "verification_status": str(row.get("verification_status", "") or row.get("candidate_status", "") or ""),
            "page_ids": page_ids,
            "page_labels": _page_labels(page_ids, data),
        })
    return rows

def serialise_chat_messages(data):
    messages = data.get("messages")
    out = []
    if isinstance(messages, pd.DataFrame) and not messages.empty:
        frame = messages.copy()
        if "created_at" in frame.columns:
            frame = frame.sort_values("created_at")
        for _, row in frame.iterrows():
            role = str(row.get("role", ""))
            if role in {"user", "assistant"}:
                out.append({"role": role, "content": str(row.get("message_text", ""))})
    return out


# =============================================================================
# ASYNC JOB COORDINATOR
#
# Jobs are registered per case and stage so any browser session (or a
# reload) can find the job that is running for a case and resume showing
# its progress. Finished jobs stay readable for FINISHED_JOB_TTL_SECONDS.
# A stage job also records its lifecycle in the workflow state: running
# when it starts, error with the message when it fails, and completed
# when it succeeds (unless the task already set a more precise outcome,
# e.g. waiting for the user).
# =============================================================================
FINISHED_JOB_TTL_SECONDS = 15 * 60


def _prune_jobs():
    now = time.time()
    with JOBS_LOCK:
        for job_id in [
            key for key, job in JOBS.items()
            if job.get("finished_at") and now - job["finished_at"] > FINISHED_JOB_TTL_SECONDS
        ]:
            JOBS.pop(job_id, None)


def _new_job(case_id="", stage=""):
    _prune_jobs()
    job_id = uuid.uuid4().hex
    try:
        owner = getattr(g, "user", "")
    except RuntimeError:                       # outside a request
        owner = ""
    with JOBS_LOCK:
        JOBS[job_id] = {
            "status": "running",
            "user": owner,
            "case_id": str(case_id or ""),
            "stage": stage,
            "progress": {"stage": stage, "phase": "", "current": 0, "total": 0, "detail": ""},
            "result": None,
            "error": None,
            "finished_at": None,
        }
    return job_id


def _set_progress(job_id, **kwargs):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        if job:
            job["progress"].update(kwargs)


def _finish_job(job_id, result=None, error=None):
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        if not job:
            return
        if error is not None:
            job["status"] = "error"
            job["error"] = error
        else:
            job["status"] = "done"
            job["result"] = result
        job["finished_at"] = time.time()


def _active_jobs(case_id):
    """{stage: {job_id, progress}} for jobs still running on this case."""
    active = {}
    with JOBS_LOCK:
        for job_id, job in JOBS.items():
            if job["status"] == "running" and job.get("case_id") == str(case_id) and job.get("stage"):
                active[job["stage"]] = {"job_id": job_id, "progress": dict(job["progress"])}
    return active


def _run_job(job_id, target, session_id=None, use_case="litigation"):
    try:
        result = target()
        _finish_job(job_id, result=result)
    except Exception as error:
        traceback.print_exc()
        if session_id:
            increment_usage(session_id, "error_count", 1, use_case=use_case)
        _finish_job(job_id, error=f"{type(error).__name__}: {error}")


def _run_async(job_id, target, session_id=None, use_case="litigation"):
    if RUN_JOBS_INLINE:
        _run_job(job_id, target, session_id=session_id, use_case=use_case)
        return

    @copy_current_request_context
    def _thread_target():
        _run_job(job_id, target, session_id=session_id, use_case=use_case)

    thread = threading.Thread(target=_thread_target, daemon=True)
    thread.start()


def _set_phase(job_id, phase, detail="", current=0, total=0):
    _set_progress(job_id, phase=phase, detail=detail, current=current, total=total)


def _start_stage_job(case_id, stage, phase, detail, task, session_id=None, use_case="litigation"):
    """Run task(job_id) in the background as the given workflow stage.

    Returns a Flask response carrying the job id. If a job for this stage
    is already running on the case, its id is returned instead of starting
    a duplicate, so a double click or a second browser tab simply follows
    the existing job.
    """
    existing = _active_jobs(case_id).get(stage)
    if existing:
        return jsonify({"job_id": existing["job_id"], "already_running": True})

    job_id = _new_job(case_id, stage)
    _set_phase(job_id, phase, detail)
    update_workflow_state(
        case_id,
        stages={stage: stage_patch(RUNNING, phase=phase, detail=detail, job_id=job_id)},
        action=f"{stage}_started",
    )

    def _wrapped():
        try:
            result = task(job_id)
        except Exception as error:
            with JOBS_LOCK:
                failed_phase = (JOBS.get(job_id) or {}).get("progress", {}).get("phase", phase)
            try:
                update_workflow_state(
                    case_id,
                    stages={stage: stage_patch(ERROR, phase=failed_phase, error=f"{type(error).__name__}: {error}")},
                    action=f"{stage}_failed",
                )
            except Exception:
                traceback.print_exc()
            raise

        record = (read_workflow_state(case_id).get("stages") or {}).get(stage) or {}
        if record.get("status") == RUNNING and record.get("job_id") == job_id:
            update_workflow_state(case_id, stages={stage: stage_patch(COMPLETED)}, action=f"{stage}_completed")
        return result

    _run_async(job_id, _wrapped, session_id=session_id, use_case=use_case)
    return jsonify({"job_id": job_id})


# =============================================================================
# SIGNED-IN USER AND CASE OWNERSHIP (every request)
#
# The user is the one signed in to Dataiku. A request about a case (its
# case_id in the query, the form or the JSON body) or about a background job
# is answered only for that case's owner; anything else gets "Case not
# found", the same as a case that does not exist.
# =============================================================================
@app.before_request
def _authorise_request():
    if request.method == "OPTIONS":
        return None
    user = access.user_from_headers(request.headers)
    g.user = user
    body = request.get_json(silent=True) if request.is_json else None
    case_id = access.requested_case_id(request.args, request.form, body)
    job_case = None
    if request.path.endswith("/job_status"):
        with JOBS_LOCK:
            job = JOBS.get(request.args.get("job_id", ""))
        if job and job.get("user") and job["user"] != user:
            return jsonify({"status": "unknown"}), 404
        job_case = (job or {}).get("case_id") or None
    status, message = access.authorise(user, case_id, job_case)
    if status != 200:
        return jsonify({"error": message}), status
    return None


@app.route("/me")
def me():
    return _ok({"user": g.user})


@app.route("/job_status")
def job_status():
    job_id = request.args.get("job_id", "")
    with JOBS_LOCK:
        job = JOBS.get(job_id)
        if not job:
            return jsonify({"status": "unknown"}), 404
        payload = {
            "status": job["status"],
            "stage": job.get("stage", ""),
            "progress": dict(job["progress"]),
            "error": job["error"],
        }
        if job["status"] == "done":
            payload["result"] = job["result"]
    return _ok(payload)


@app.route("/diagnostics")
def diagnostics():
    case_id = request.args.get("case_id", "")
    if not case_id:
        return jsonify({"error": "case_id is required"}), 400
    data = load_case_data(case_id)
    counts = {}
    for key, frame in data.items():
        if isinstance(frame, pd.DataFrame):
            counts[key] = int(len(frame))
        elif isinstance(frame, dict):
            counts[key] = len(frame)
    state = restore_workflow_state(data)
    return _ok({
        "case_id": case_id,
        "row_counts": counts,
        "extraction_reached_facts": bool(counts.get("facts", 0) or counts.get("fact_candidates", 0)),
        "workflow": {
            "has_attorney_summary": bool(state.get("attorney_summary")),
            "summary_approved": bool(state.get("summary_approved")),
            "has_analysis": bool(state.get("analysis")),
            "has_memo": bool(state.get("memo")),
            "case_dirty": bool(state.get("case_dirty")),
            "dirty_reason": state.get("dirty_reason", ""),
            "stages": state.get("stages") or {},
        },
        "run_jobs_inline": RUN_JOBS_INLINE,
    })


# =============================================================================
# FLASK APPLICATION ROUTES
# =============================================================================
@app.route("/bootstrap")
def bootstrap():
    return jsonify({
        "logo": _load_app_logo_data_uri(),
        "agreement_types": list(AGREEMENT_TYPES),
        "relationship_types": list(RELATIONSHIP_TYPES),
    })


@app.route("/cases")
def cases_endpoint():
    """The signed-in user's own cases only (others' are never listed)."""
    workflow = request.args.get("workflow", "litigation")
    query = request.args.get("query", "")
    show_archived = str(request.args.get("archived", "")).lower() in {"1", "true", "yes"}
    records = {r["case_id"]: r for r in access.user_cases(g.user)}
    cases = list_cases()
    if not cases.empty and "case_id" in cases.columns:
        cases = cases[cases["case_id"].astype(str).isin(records)]
        keep = cases["case_id"].astype(str).map(lambda cid: bool(records[cid].get("archived")) == show_archived)
        cases = cases[keep]
    filtered = filter_cases_by_workflow(cases, workflow)
    filtered = case_search_frame(filtered, query)
    if "updated_at" in filtered.columns:
        filtered = filtered.sort_values("updated_at", ascending=False)
    cards = []
    for row in frame_to_records(filtered):
        card = serialise_case_card(row, workflow)
        record = records.get(card["case_id"]) or {}
        card["details"] = record.get("details") or {}
        card["archived"] = bool(record.get("archived"))
        card["progress"] = _case_progress(card["case_id"], workflow)
        cards.append(card)
    return _ok({"cases": cards})


def _case_progress(case_id, workflow):
    """Where a case stands, from its saved workflow state only (cheap):
    new / documents_uploaded / processing / accounting_review /
    legal_analysis / ready_for_pleading / attorney_review / completed."""
    if workflow == "agreement_review":
        return ""
    state = read_workflow_state(case_id)
    stages = state.get("stages") or {}
    accounting_done = str(state.get("accounting_status") or "") in {"forensic_complete", "complete_no_transactions"}
    if state.get("memo"):
        return "completed" if state.get("pleading_status") == "final" else "attorney_review"
    if any((stage or {}).get("status") == RUNNING for stage in stages.values()):
        return "processing"
    if state.get("analysis") and accounting_done:
        return "ready_for_pleading"
    if str(state.get("accounting_status") or "") not in {"", "not_started"} and not accounting_done:
        return "accounting_review"
    if (stages.get("documents") or {}).get("status") == COMPLETED:
        return "legal_analysis" if state.get("attorney_summary") or accounting_done else "documents_uploaded"
    return "new"


@app.route("/case/archive", methods=["POST"])
def case_archive():
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    access.update_case_record(case_id, g.user, {"archived": bool(body.get("archived", True))})
    return _ok({"archived": bool(body.get("archived", True))})


@app.route("/create_case", methods=["POST"])
def create_case_endpoint():
    body = request.get_json(force=True)
    case_name = str(body.get("case_name", "")).strip()
    language = body.get("language", "en")
    workflow_type = body.get("workflow_type", "litigation")
    if not case_name:
        return jsonify({"error": "case_name is required"}), 400
    details = {key: body.get(key, "") for key in access.DETAIL_FIELDS}
    case_id = create_case(
        case_name=case_name,
        intake_mode="mixed",
        language=language,
        workflow_type=workflow_type,
        created_by=g.user,
        description=str(body.get("description", "") or ""),
        client_name=str(body.get("client_name", "") or ""),
        matter_type=str(body.get("case_type", "") or body.get("contract_type", "") or ""),
    )
    access.register_case(case_id, g.user, workflow_type, details)
    return jsonify({"case_id": case_id})



@app.route("/case")
def case_endpoint():
    case_id = request.args.get("case_id", "")
    if not case_id:
        return jsonify({"error": "case_id is required"}), 400
    data = load_case_data(case_id)
    case = data["case"]
    workflow_type = str(case.get("workflow_type", "") or "").strip() or "litigation"

    response = {
        "case_id": case_id,
        "reference": short_case_reference(case_id),
        "workflow_type": workflow_type,
        "case": {k: _clean_scalar(v) for k, v in case.items()},
        "display_name": first_case_value(case, "client_name", "customer_name", "case_name", default=""),
        "counts": {
            "documents": int(len(data["documents"])),
            "pages": int(len(data["pages"])),
            "facts": int(len(best(data, "facts", "fact_candidates"))),
            "parties": int(len(data["parties"])),
            "issues": int(len(best(data, "issues", "issue_candidates"))),
        },
    }

    if workflow_type == "agreement_review":
        stored = load_agreement_states(case_id)
        response["agreement_state"] = {
            "classification": stored.get("classification"),
            "profile": stored.get("confirmed_profile"),
            "clause_map": stored.get("clause_map"),
            "authorities": stored.get("authorities"),
            "review": stored.get("review"),
        }
        return _ok(response)

    state = restore_workflow_state(data)
    corrections = load_latest_corrections(case_id)
    resolutions = load_fact_resolutions(case_id)
    review_items = _review_items(case_id, data, corrections, resolutions)
    accounting_meta = accounting_state(data, state, review_items)
    register = case_register(data)
    overview = stage_overview(case_id, state, data, accounting_meta, register)

    response["workflow_state"] = state
    response["stages"] = overview["stages"]
    response["waiting"] = overview["waiting"]
    response["running"] = overview["running"]
    response["errors"] = overview["errors"]
    response["next_stage"] = overview["next_stage"]
    # Legacy names kept for callers that still read them.
    response["steps"] = overview["stages"]
    response["next_action_key"] = overview["next_stage"]

    response["chat_messages"] = serialise_chat_messages(data)
    response["case_register"] = register
    response["facts_register"] = serialise_facts_register(data)
    response["summary_support"] = serialise_summary_support(state.get("attorney_summary") or {}, data)

    ledger, withheld = _normalized_ledger(case_id, data["financial_line_items"], corrections, resolutions)
    classifications = _latest_classifications_by_page(data)
    accounting_record = (state.get("stages") or {}).get("accounting") or {}
    response["accounting"] = {
        "status": accounting_meta["status"],
        "dirty": accounting_meta["dirty"],
        "dirty_reason": accounting_meta["dirty_reason"],
        "has_classifications": accounting_meta["has_classifications"],
        "has_line_items": accounting_meta["has_line_items"],
        "has_pending_review": accounting_meta["has_pending_review"],
        "pending_count": accounting_meta["pending_count"],
        "pending_field_count": accounting_meta["pending_field_count"],
        "documents": [
            {
                "case_document_id": str(row.get("case_document_id", "")),
                "original_filename": row.get("original_filename") or row.get("case_document_id", ""),
                "classification": {},
            }
            for row in frame_to_records(data["documents"])
        ],
        "page_classification_counts": _classification_counts(classifications),
        "last_run_report": accounting_record.get("report") or {},
        "review_items": review_items,
        "pending_conflicts": review_items,
        "normalized_ledger": ledger,
        "withheld_rows": withheld,
        "instructions": state.get("accounting_instructions", ""),
        "run_metadata": data.get("forensic_run_metadata") or {},
        "cross_check_summary": data.get("cross_check_summary") or {},
        "discrepancies": data.get("discrepancies") or [],
        "findings": _claims_with_evidence(data.get("forensic_findings"), ledger, data),
        "facts_table": _fact_table(case_id, data, corrections, resolutions),
        "document_claims": _document_claims(case_id, data),
        "financial_pages": _financial_pages(case_id, data, classifications),
        "facts_confirmed": bool(state.get("accounting_facts_confirmed")),
        "facts_confirmed_by": state.get("accounting_facts_confirmed_by", ""),
        "facts_confirmed_at": state.get("accounting_facts_confirmed_at", ""),
    }
    return _ok(response)


def _document_claims(case_id, data):
    """Financial claims found in the accounting pages (what an email or a
    narrative section SAYS about money): to be verified, never evidence."""
    labels = _page_reference_map(data)
    out = []
    for page_id, claims in (load_document_claims(case_id) or {}).items():
        for claim in claims or []:
            out.append({**claim, "page_id": page_id, "page_label": labels.get(page_id, "")})
    out.sort(key=lambda c: (str(c.get("case_document_id", "")), int(float(c.get("page_number") or 0))))
    return out


def _fact_table(case_id, data, corrections=None, resolutions=None):
    """Every atomic fact for the review table (ledger, uncertain, deleted)."""
    line_items = data.get("financial_line_items")
    if not isinstance(line_items, pd.DataFrame) or line_items.empty:
        return []
    if corrections is None:
        corrections = load_latest_corrections(case_id)
    if resolutions is None:
        resolutions = load_fact_resolutions(case_id)
    return build_fact_table(line_items.to_dict(orient="records"), corrections, resolutions)


def _financial_pages(case_id, data, classifications=None):
    """The pages shown beside the fact table: pages classified financial /
    mixed and pages that have facts, in document and page order."""
    pages = data.get("pages")
    if not isinstance(pages, pd.DataFrame) or pages.empty:
        return []
    if classifications is None:
        classifications = _latest_classifications_by_page(data)
    with_facts = set()
    line_items = data.get("financial_line_items")
    if isinstance(line_items, pd.DataFrame) and not line_items.empty and "page_id" in line_items.columns:
        with_facts = set(line_items["page_id"].astype(str))
    documents = {str(r.get("case_document_id", "")): r for r in frame_to_records(data["documents"])}
    document_order = list(documents)
    labels = _page_reference_map(data)
    id_col = _page_id_column(pages)
    out = []
    for row in frame_to_records(pages):
        page_id = str(row.get(id_col, "") or row.get("page_id", ""))
        info = classifications.get(page_id, {})
        financial = info.get("page_type") in {"financial", "mixed"} and info.get("status") != "failed"
        if not (financial or page_id in with_facts):
            continue
        document_id = str(row.get("case_document_id", ""))
        try:
            number = int(float(row.get("page_number") or 0))
        except (TypeError, ValueError):
            number = 0
        out.append({
            "page_id": page_id,
            "page_number": number,
            "case_document_id": document_id,
            "document_name": (documents.get(document_id) or {}).get("original_filename") or document_id,
            "label": labels.get(page_id, ""),
            "page_type": info.get("page_type", ""),
            "image_url": f"/page_image?case_id={case_id}&page_id={page_id}",
            "_order": (document_order.index(document_id) if document_id in document_order else 999, number),
        })
    unique = {item["page_id"]: item for item in out}
    ordered = sorted(unique.values(), key=lambda item: item["_order"])
    for item in ordered:
        item.pop("_order", None)
    return ordered


CLAIM_RESULT_ORDER = (
    "CONTRADICTED", "PARTIALLY_SUPPORTED", "SUPPORTED",
    "NOT_VERIFIABLE", "NO_FINANCIAL_EVIDENCE", "NOT_FINANCIAL_CLAIM",
)


EVIDENCE_GROUPS = (
    "supporting_evidence",
    "partially_supporting_evidence",
    "contradicting_evidence",
    "unresolved_evidence",
)


def _claims_with_evidence(findings, ledger, data):
    """Claim evaluations with every cited ledger fact resolved to its
    values, status and source page, so each accounting position can be
    checked against the document it relies on.

    For supporting and partial evidence the total of the cited facts'
    values is computed here (per currency), independently of the model's
    substantiated_amount. A cited id no longer in the ledger (cleared or
    re-extracted after the analysis ran) is listed rather than dropped."""
    by_id = {str(entry.get("row_id")): entry for entry in ledger}
    labels = _page_reference_map(data)
    fields = ("fact_type", "date", "description", "value", "value_field", "currency", "debit", "credit",
              "balance", "transaction_reference", "counterparty", "status", "page_id")

    def resolve(ids, stale):
        rows = []
        for record_id in ids or []:
            entry = by_id.get(str(record_id))
            if entry is None:
                stale.append(str(record_id))
                continue
            row = {key: entry.get(key) for key in fields}
            row["row_id"] = entry.get("row_id")
            row["page_label"] = labels.get(str(entry.get("page_id", "")), "")
            # Names the earlier view used.
            row["amount"] = entry.get("value")
            row["reference_number"] = entry.get("transaction_reference") or ""
            row["debit_or_credit"] = entry.get("debit_or_credit", "")
            rows.append(row)
        return rows

    def totals(groups):
        sums = {}
        seen = set()
        for group in groups:
            for row in group:
                if row["row_id"] in seen:
                    continue
                seen.add(row["row_id"])
                value = parse_amount(row.get("value"))
                if value is not None:
                    currency = row.get("currency") or "—"
                    sums[currency] = sums.get(currency, 0) + value
        return {currency: format_amount(total) for currency, total in sums.items()}

    evaluations, counts = [], {}
    for raw in (findings or {}).get("claim_evaluations") or []:
        if not isinstance(raw, dict):
            continue
        item = dict(raw)
        stale = []
        for group in EVIDENCE_GROUPS:
            entries = []
            for entry in item.get(group) or []:
                if isinstance(entry, dict):
                    entries.append({
                        "explanation": entry.get("explanation", ""),
                        "record_ids": entry.get("record_ids") or [],
                        "rows": resolve(entry.get("record_ids"), stale),
                    })
            item[group] = entries
        # Results saved before evidence was grouped.
        if not any(item[group] for group in EVIDENCE_GROUPS):
            if item.get("evidence_record_ids"):
                item["supporting_evidence"] = [{"explanation": "", "record_ids": item["evidence_record_ids"],
                                                "rows": resolve(item["evidence_record_ids"], stale)}]
            for entry in item.get("contradictions") or []:
                if isinstance(entry, dict):
                    item["contradicting_evidence"].append({
                        "explanation": entry.get("description", ""), "record_ids": entry.get("record_ids") or [],
                        "rows": resolve(entry.get("record_ids"), stale)})
        item["evidence_rows"] = [row for entry in item["supporting_evidence"] + item["partially_supporting_evidence"]
                                 for row in entry["rows"]]
        item["evidence_totals"] = totals([entry["rows"] for entry in
                                          item["supporting_evidence"] + item["partially_supporting_evidence"]])
        item["missing_evidence"] = [str(x) for x in (item.get("missing_evidence") or []) if str(x).strip()]
        item["accounting_position"] = item.get("accounting_position") or item.get("accounting_response") or ""
        item["no_longer_in_ledger"] = list(dict.fromkeys(stale))
        result = str(item.get("result") or "NOT_VERIFIABLE").upper()
        counts[result] = counts.get(result, 0) + 1
        evaluations.append(item)

    rank = {result: position for position, result in enumerate(CLAIM_RESULT_ORDER)}
    evaluations.sort(key=lambda item: rank.get(str(item.get("result") or "").upper(), len(rank)))
    return {"claim_evaluations": evaluations, "result_counts": counts}


def _classification_counts(classifications):
    counts = {}
    for item in classifications.values():
        key = "failed" if item.get("status") == "failed" else (item.get("page_type") or "other")
        counts[key] = counts.get(key, 0) + 1
    return counts


def _load_state(case_id):
    data = load_case_data(case_id)
    return data, restore_workflow_state(data)


@app.route("/state/persist", methods=["POST"])
def state_persist():
    """Merge the given workflow keys into the stored state. Only keys the
    caller sends are touched; stage records are managed by the server."""
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    state = body.get("state", {})
    if not case_id or not isinstance(state, dict):
        return jsonify({"error": "case_id and a state object are required"}), 400
    changes = {key: value for key, value in state.items() if key != "stages"}
    update_workflow_state(case_id, changes=changes, action="ui_state_saved")
    return _ok({"persisted": True})


# =============================================================================
# INGESTION & DOCUMENT PROCESSING
# =============================================================================
class _MemoryUpload:
    def __init__(self, name, raw, content_type=""):
        self.name = name
        self.type = content_type or mimetypes.guess_type(name or "")[0] or "application/octet-stream"
        self.size = len(raw or b"")
        self.id = uuid.uuid4().hex
        self.file_id = self.id
        self._raw = raw
        self._buffer = io.BytesIO(raw)

    def read(self, *args, **kwargs): return self._buffer.read(*args, **kwargs)
    def seek(self, *args, **kwargs): return self._buffer.seek(*args, **kwargs)
    def tell(self): return self._buffer.tell()
    def getvalue(self): return self._raw
    def getbuffer(self): return self._buffer.getbuffer()
    def close(self): self._buffer.close()


def _intake_context(case_id):
    """What the page reader may use as a hint: the case record and the
    parties already known (from documents processed earlier)."""
    case = latest_case(case_id) or {}
    parties = []
    for row in frame_to_records(case_rows("case_parties", case_id))[:20]:
        name = str(row.get("party_name") or row.get("name") or "").strip()
        role = str(row.get("party_role") or row.get("role") or "").strip()
        if name and name not in [p.split(" (")[0] for p in parties]:
            parties.append(f"{name} ({role})" if role else name)
    return {"case": {k: _clean_scalar(v) for k, v in case.items()}, "parties": parties}


@app.route("/documents/process", methods=["POST"])
def documents_process():
    """Stage "documents": pages -> page pipeline (PDF text or VLM reading,
    consolidated) -> page summaries ->
    parties / events / facts / issues / evidence merged into the case
    register. New material marks downstream work stale."""
    case_id = request.form.get("case_id", "")
    file_purpose = request.form.get("file_purpose", "full_case")
    session_id = request.form.get("session_id")
    render_zoom = 1.6
    uploaded = request.files.getlist("files")
    payloads = [(f.filename, f.read(), f.mimetype) for f in uploaded]
    if not case_id or not payloads:
        return jsonify({"error": "case_id and at least one file are required"}), 400

    def task(job_id):
        total_pages, usable_pages = 0, 0
        totals = {"facts": 0, "issues": 0, "parties": 0, "events": 0, "evidence_requests": 0}
        status_counts, failure_notes = {}, []
        llm_call_count = 0

        for file_index, (name, raw, content_type) in enumerate(payloads, start=1):
            _set_phase(job_id, "extracting_pages", f"{name} ({file_index}/{len(payloads)}): rendering pages…")
            document_row, pdf_bytes = save_upload(
                case_id, _MemoryUpload(name, raw, content_type), document_type=file_purpose
            )
            def update_progress(current, total, page_number, state, _name=name):
                _set_phase(job_id, "extracting_pages",
                           f"{_name}: page {page_number}/{total} — {state}", current=current, total=total)
            pages = extract_pdf_page_by_page(
                case_id=case_id,
                case_document_id=document_row["case_document_id"],
                pdf_bytes=pdf_bytes,
                zoom=render_zoom,
                progress_callback=update_progress,
                document_name=name,
                case_context=_intake_context(case_id),
            )
            llm_call_count += len(pages)
            completed_pages = [p for p in pages if p.get("processing_status") in {"completed", "completed_review_required"}]
            for page in pages:
                page_status = str(page.get("processing_status", "") or "unknown")
                status_counts[page_status] = status_counts.get(page_status, 0) + 1
                if page_status in {"completed", "completed_review_required"}:
                    continue
                for note_key in ("processing_error", "error", "processing_note", "note", "failure_reason"):
                    if page.get(note_key):
                        note = str(page[note_key])[:300]
                        if note not in failure_notes:
                            failure_notes.append(note)
                        break

            total_pages += len(pages)
            usable_pages += len(completed_pages)
            if completed_pages:
                _set_phase(job_id, "building_case_register", f"{name}: extracting parties, events, facts and issues…")
                case_map = build_case_map(completed_pages)
                llm_call_count += 1
                counts = persist_case_map(case_id, case_map)
                for key in totals:
                    totals[key] += int(counts.get(key, 0) or 0)

        if llm_call_count:
            increment_usage(session_id, "llm_request_count", llm_call_count)
        increment_usage(session_id, "attachment_count", len(payloads))

        report = {
            "files": len(payloads),
            "usable": usable_pages,
            "total": total_pages,
            "page_status": status_counts,
            "page_errors": failure_notes[:3],
            **totals,
        }
        failed = total_pages - usable_pages
        update_workflow_state(
            case_id,
            changes=invalidation_changes("New document or evidence was uploaded."),
            stages={"documents": stage_patch(
                COMPLETED,
                detail=f"{usable_pages}/{total_pages} pages usable" + (f", {failed} failed" if failed else ""),
                report=report,
            )},
            action="documents_processed",
        )
        return report

    return _start_stage_job(
        case_id, "documents", "extracting_pages", "Preparing documents…", task, session_id=session_id,
    )


@app.route("/documents/review_completeness", methods=["POST"])
def documents_review_completeness():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    session_id = body.get("session_id")
    data, state = _load_state(case_id)
    interview_state = body.get("interview_state")
    result = assess_next_step(
        case_record=data["case"], messages=data["messages"], documents=data["documents"], pages=data["pages"],
        facts=data["facts"], fact_candidates=data["fact_candidates"], parties=data["parties"],
        issues=data["issues"], issue_candidates=data["issue_candidates"], evidence=data["evidence"],
        legal_research=((state.get("research") or {}).get("authority_nodes", [])),
        previous_state=interview_state,
    )
    increment_usage(session_id, "llm_request_count", 1)
    increment_usage(session_id, "message_count", 1)
    persist_interview_state(case_id, result["decision"].payload)
    return _ok({"reply": result["reply"], "interview_state": result["decision"].payload})


# ============================================================
# Accounting & forensic dispute analysis — stage "accounting"
#
#   1. classify pages (financial / claim / mixed / other). Stored results
#      are reused; only new pages and failed attempts are classified.
#   2. extract financial rows from financial + mixed pages only, reusing
#      stored rows (PyMuPDF structure + intake OCR -> LLM reconstruction).
#   3. every uncertain or missing required value becomes a pending review
#      item; the stage WAITS for the user until all are answered.
#   4. corrections are stored separately (never overwrite the extraction);
#      the normalized ledger contains only fully reviewed rows.
#   5. claims vs evidence analysis runs only on the reviewed ledger, with
#      the user's instructions stored alongside the result.
# ============================================================
def _selected_page_ids(data, document_ids=None):
    pages = data["pages"]
    if pages.empty:
        return []
    id_col = _page_id_column(pages)
    if document_ids:
        doc_col = "case_document_id" if "case_document_id" in pages.columns else "document_id"
        pages = pages[pages[doc_col].astype(str).isin([str(d) for d in document_ids])]
    return pages[id_col].dropna().astype(str).tolist()


def _accounting_pipeline(job_id, case_id, document_ids=None, extract=True, session_id=None):
    data = load_case_data(case_id)
    before = restore_workflow_state(data)
    page_ids = _selected_page_ids(data, document_ids)
    if not page_ids:
        raise ValueError("No processed document pages are available. Process the case documents first.")

    # 1. Classification — only pages without a usable stored result.
    to_classify = pages_needing_classification(case_id, page_ids)
    if to_classify:
        _set_phase(job_id, "classifying_pages", f"Classifying {len(to_classify)} page(s) as financial / non-financial…")
        classify_case_pages(case_id, to_classify)
        if session_id:
            increment_usage(session_id, "llm_request_count", len(to_classify))

    classifications = _latest_classifications_by_page(load_case_data(case_id))
    in_scope = [classifications.get(pid, {}) for pid in page_ids]
    classification_failures = sum(1 for item in in_scope if item.get("status") == "failed")
    financial_page_ids = [
        pid for pid in page_ids
        if classifications.get(pid, {}).get("page_type") in {"financial", "mixed"}
        and classifications.get(pid, {}).get("status") != "failed"
    ]
    report = {
        "pages_in_scope": len(page_ids),
        "pages_classified_now": len(to_classify),
        "classification_failures": classification_failures,
        "financial_pages": len(financial_page_ids),
    }

    if not extract:
        update_workflow_state(
            case_id,
            changes={"accounting_status": "classified"} if before.get("accounting_status") in {"", "not_started"} else {},
            stages={"accounting": stage_patch(READY, detail="pages_classified", report=report)},
            action="accounting_classified",
        )
        return report

    # 2. Extraction — only financial pages not yet extracted.
    extraction = {"new_rows_appended": 0, "rows_discarded": 0, "page_failures": [], "pages_processed": 0}
    if financial_page_ids:
        _set_phase(job_id, "extracting_financial_data",
                   f"Extracting financial data from {len(financial_page_ids)} page(s)…", total=len(financial_page_ids))

        def _on_progress(current, total, page_id):
            _set_phase(job_id, "extracting_financial_data",
                       f"Extracting financial data (page {current} of {total})…", current=current, total=total)

        try:
            extraction = run_financial_extraction(
                case_id, financial_page_ids, progress_callback=_on_progress,
                page_types={pid: classifications.get(pid, {}).get("page_type", "financial") for pid in financial_page_ids},
            )
        except ValueError as error:
            # No page text available for the selected pages.
            report["extraction_note"] = str(error)
        if session_id:
            increment_usage(session_id, "llm_request_count", int(extraction.get("pages_processed", 0) or 0))

    report.update({
        "pages_extracted_now": int(extraction.get("pages_processed", 0) or 0),
        "rows_added": int(extraction.get("new_rows_appended", 0) or 0),
        "rows_discarded": int(extraction.get("rows_discarded", 0) or 0),
        "extraction_failures": len(extraction.get("page_failures") or []),
    })

    # 3. Decide where accounting stands now.
    _set_phase(job_id, "checking_review_items", "Checking for values that need your confirmation…")
    final_data = load_case_data(case_id)
    pending = _review_items(case_id, final_data)
    has_items = not final_data["financial_line_items"].empty
    previous = str(before.get("accounting_status", "") or "")
    report["pending_review_items"] = sum(len(item.get("fields") or []) for item in pending)

    changes = {"accounting_dirty": False, "accounting_dirty_reason": ""}
    confirmed = bool(before.get("accounting_facts_confirmed"))
    if report["rows_added"]:
        # New facts: the user reviews the table again before the analysis.
        changes.update({"accounting_facts_confirmed": False, "accounting_facts_confirmed_by": "",
                        "accounting_facts_confirmed_at": ""})
        confirmed = False
    if pending or (has_items and not confirmed):
        changes["accounting_status"] = "needs_review"
        stage = stage_patch(WAITING, detail="review_items" if pending else "confirm_facts", report=report)
    elif (
        report["rows_added"] == 0
        and previous in {"forensic_complete", "complete_no_transactions"}
        and not before.get("accounting_dirty")
    ):
        # Nothing new: keep the completed analysis instead of forcing a rerun.
        changes = {}
        stage = stage_patch(COMPLETED, detail=previous, report=report)
    elif has_items:
        changes["accounting_status"] = "ready_for_synthesis"
        stage = stage_patch(READY, detail="ready_for_analysis", report=report)
    else:
        changes["accounting_status"] = "complete_no_transactions"
        stage = stage_patch(COMPLETED, detail="complete_no_transactions", report=report)

    update_workflow_state(case_id, changes=changes, stages={"accounting": stage}, action="accounting_pipeline_complete")
    report["accounting_status"] = changes.get("accounting_status", previous)
    report["rows_persisted"] = report["rows_added"]
    return report


@app.route("/accounting/run_auto_pipeline", methods=["POST"])
def accounting_run_auto_pipeline():
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", "")).strip()
    document_ids = [str(d).strip() for d in (body.get("document_ids") or []) if str(d).strip()]
    session_id = body.get("session_id")
    if not case_id:
        return jsonify({"error": "case_id is required"}), 400
    return _start_stage_job(
        case_id, "accounting", "classifying_pages", "Starting accounting extraction…",
        lambda job_id: _accounting_pipeline(job_id, case_id, document_ids or None, session_id=session_id),
        session_id=session_id,
    )


@app.route("/accounting/classify_documents", methods=["POST"])
def accounting_classify_documents():
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", "")).strip()
    session_id = body.get("session_id")
    if not case_id:
        return jsonify({"error": "case_id is required"}), 400
    return _start_stage_job(
        case_id, "accounting", "classifying_pages", "Classifying pages…",
        lambda job_id: _accounting_pipeline(job_id, case_id, None, extract=False, session_id=session_id),
        session_id=session_id,
    )


@app.route("/accounting/extract", methods=["POST"])
def accounting_extract():
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", "")).strip()
    document_ids = [str(d).strip() for d in (body.get("document_ids") or []) if str(d).strip()]
    session_id = body.get("session_id")
    if not case_id:
        return jsonify({"error": "case_id is required"}), 400
    if not document_ids:
        return jsonify({"error": "At least one document must be selected."}), 400
    return _start_stage_job(
        case_id, "accounting", "classifying_pages", "Starting accounting extraction…",
        lambda job_id: _accounting_pipeline(job_id, case_id, document_ids, session_id=session_id),
        session_id=session_id,
    )


# What "Clear accounting data" deletes, for this case only. Documents,
# OCR text, the case register, the attorney review and the legal analysis
# are NOT touched. A pleading built on the cleared accounting is kept but
# shown as stale until accounting is completed again.
ACCOUNTING_CLEAR_DATASETS = {
    "page_classifications": FINANCIAL_DOCUMENT_CLASSIFICATION_DATASET,
    "extracted_line_items": FINANCIAL_LINE_ITEMS_DATASET,
    "user_corrections": FINANCIAL_LINE_ITEM_CORRECTIONS_DATASET,
    "financial_timeline": FINANCIAL_TIMELINE_DATASET,
    "discrepancies": FINANCIAL_DISCREPANCIES_DATASET,
    "claim_findings": FINANCIAL_FINDINGS_DATASET,
}


@app.route("/accounting/clear", methods=["POST"])
def accounting_clear():
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    if not case_id:
        return jsonify({"error": "case_id is required"}), 400
    if _active_jobs(case_id).get("accounting"):
        return jsonify({"error": "The accounting analysis is running. Wait for it to finish before clearing."}), 409

    deleted = {}
    try:
        deleted["financial_claims"] = sum(len(c or []) for c in load_document_claims(case_id).values())
        save_document_claims(case_id, {}, replace=True)
    except Exception:
        traceback.print_exc()
        deleted["financial_claims"] = "error"
    for label, dataset_name in ACCOUNTING_CLEAR_DATASETS.items():
        try:
            before = len(case_rows(dataset_name, case_id))
            replace_case_rows(dataset_name, case_id, [])
            deleted[label] = before
        except Exception:
            traceback.print_exc()
            deleted[label] = "error"

    update_workflow_state(
        case_id,
        changes={
            "accounting_status": "not_started",
            "accounting_dirty": False,
            "accounting_dirty_reason": "",
            "accounting_facts_confirmed": False,
            "accounting_facts_confirmed_by": "",
            "accounting_facts_confirmed_at": "",
        },
        stages={"accounting": stage_patch(NOT_STARTED, finished_at="", report={})},
        action="accounting_cleared",
    )
    return _ok({"cleared": True, "deleted": deleted})


def _validate_correction(field_name, value):
    """(clean_value, error). Only values that downstream analysis can use
    are accepted, so a correction can never re-introduce an unusable field."""
    value = str(value or "").strip()
    if field_name not in FIELD_NAMES:
        return None, f"Unknown field '{field_name}'."
    if not value:
        return None, "A value is required."
    if field_name == "debit_or_credit":
        direction = normalize_debit_credit(value, value)
        if direction not in {"debit", "credit"}:
            return None, "Enter 'debit' or 'credit'."
        return direction, None
    if field_name in {"amount", "running_balance"}:
        if not re.fullmatch(r"-?[\d,]+\.\d{2}", normalize_amount(value)):
            return None, "Enter a numeric amount, e.g. 12,450.00."
    return value, None


@app.route("/accounting/correction", methods=["POST"])
def accounting_correction():
    """Store the user's answer for one pending field. When the last pending
    item is answered, accounting is ready to continue (ready_for_analysis)
    and the response says so, with the instructions last used."""
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    row_id = str(body.get("row_id", ""))
    field_name = str(body.get("field_name", ""))
    corrected_by = str(body.get("corrected_by", "") or "attorney")
    note = str(body.get("note", "") or "")
    if not (case_id and row_id and field_name):
        return jsonify({"error": "case_id, row_id and field_name are required."}), 400
    value, error = _validate_correction(field_name, body.get("value", ""))
    if error:
        return jsonify({"error": error}), 400

    submit_correction(case_id, row_id, field_name, value, corrected_by=corrected_by, correction_note=note)
    result = _after_review_change(case_id, corrected_by)
    result["value"] = value
    return _ok(result)


def _after_review_change(case_id, actor):
    """Move the accounting stage on after the user answered a review item:
    still waiting, ready to continue (last item answered), or — when the
    analysis had already run — ready again with the old results marked
    out of date."""
    rows = case_rows(FINANCIAL_LINE_ITEMS_DATASET, case_id)
    pending = count_uncertain(
        rows.to_dict(orient="records") if not rows.empty else [],
        load_latest_corrections(case_id),
        load_fact_resolutions(case_id),
    )
    state = read_workflow_state(case_id)
    status = str(state.get("accounting_status", "") or "")
    confirmed = bool(state.get("accounting_facts_confirmed"))

    if pending or not confirmed:
        update_workflow_state(
            case_id,
            changes={"accounting_status": "needs_review"},
            stages={"accounting": stage_patch(WAITING, detail="review_items" if pending else "confirm_facts")},
            action="accounting_fact_reviewed",
            actor=actor,
        )
    elif status in {"forensic_complete", "complete_no_transactions"}:
        update_workflow_state(
            case_id,
            changes={
                "accounting_status": "ready_for_synthesis",
                "accounting_dirty": True,
                "accounting_dirty_reason": "The ledger was changed after the accounting analysis ran.",
            },
            stages={"accounting": stage_patch(READY, detail="ready_for_analysis")},
            action="accounting_ledger_changed",
            actor=actor,
        )
    else:
        update_workflow_state(
            case_id,
            changes={"accounting_status": "ready_for_synthesis"},
            stages={"accounting": stage_patch(READY, detail="ready_for_analysis")},
            action="accounting_review_resolved",
            actor=actor,
        )

    return {
        "saved": True,
        "pending_count": pending,
        "facts_confirmed": confirmed,
        "ready_for_analysis": not pending and confirmed,
        "instructions": state.get("accounting_instructions", ""),
    }


def _fact_row(case_id, row_id):
    frame = case_rows(FINANCIAL_LINE_ITEMS_DATASET, case_id)
    if frame.empty or "row_id" not in frame.columns:
        return None
    match = frame[frame["row_id"].astype(str) == str(row_id)]
    return match.iloc[-1].to_dict() if not match.empty else None


@app.route("/accounting/fact/resolve", methods=["POST"])
def accounting_fact_resolve():
    """The user's decision on one uncertain fact:
      confirm         accept the fact as the AI interpreted it
      correct         replace fields with the user's values (validated)
      ignore          leave the fact out of the ledger
      accept_proposal accept the rewrite produced from the user's explanation
    The extracted fact is never overwritten; decisions are layered on top."""
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    row_id = str(body.get("row_id", ""))
    action = str(body.get("action", ""))
    actor = str(body.get("decided_by", "") or "attorney")
    explanation = str(body.get("explanation", "") or "").strip()
    if not (case_id and row_id) or action not in {"confirm", "correct", "ignore", "accept_proposal"}:
        return jsonify({"error": "case_id, row_id and a valid action are required."}), 400

    row = _fact_row(case_id, row_id)
    if row is None:
        return jsonify({"error": "That fact no longer exists. Refresh the case."}), 404

    if action == "correct":
        fields, errors = validate_fact_fields(body.get("fields") or {})
        if errors:
            return jsonify({"error": "; ".join(errors)}), 400
        if not fields:
            return jsonify({"error": "Enter at least one corrected value."}), 400
        submit_fact_resolution(case_id, row_id, "correct", fields, explanation, actor)
    elif action == "accept_proposal":
        current = effective_fact(row, load_latest_corrections(case_id), load_fact_resolutions(case_id))
        proposal = current.get("proposal")
        if not proposal:
            return jsonify({"error": "There is no proposal to accept for this fact."}), 409
        submit_fact_resolution(case_id, row_id, "correct", proposal.get("fields") or {},
                               proposal.get("explanation", ""), actor, proposal.get("summary", ""))
    else:
        submit_fact_resolution(case_id, row_id, action, {}, explanation, actor)

    return _ok(_after_review_change(case_id, actor))


@app.route("/accounting/fact/resolve_bulk", methods=["POST"])
def accounting_fact_resolve_bulk():
    """Confirm or ignore several uncertain facts in one request (one
    dataset write), e.g. "Confirm all on this page". Facts that are no
    longer uncertain are skipped, so a repeated click is harmless."""
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    action = str(body.get("action", ""))
    actor = str(body.get("decided_by", "") or "attorney")
    row_ids = [str(r) for r in (body.get("row_ids") or []) if str(r).strip()]
    if not case_id or action not in {"confirm", "ignore"} or not row_ids:
        return jsonify({"error": "case_id, row_ids and action (confirm or ignore) are required."}), 400

    frame = case_rows(FINANCIAL_LINE_ITEMS_DATASET, case_id)
    rows = {str(r.get("row_id")): r for r in (frame.to_dict(orient="records") if not frame.empty else [])}
    corrections, resolutions = load_latest_corrections(case_id), load_fact_resolutions(case_id)
    wanted = [
        row_id for row_id in dict.fromkeys(row_ids)
        if row_id in rows and effective_fact(rows[row_id], corrections, resolutions)["status"] == UNCERTAIN
    ]
    if wanted:
        submit_fact_resolutions(case_id, [{"row_id": row_id, "action": action} for row_id in wanted], actor)
    result = _after_review_change(case_id, actor)
    result["applied"] = len(wanted)
    return _ok(result)


@app.route("/accounting/page_extraction")
def accounting_page_extraction():
    """What the system extracted from one page before facts were made: the
    consolidated text (markdown tables), uncertain values, the route and
    image kinds, and the candidate names (from the stored extraction
    record; pages processed before it existed show their page text)."""
    case_id = request.args.get("case_id", "")
    page_id = request.args.get("page_id", "")
    if not (case_id and page_id):
        return jsonify({"error": "case_id and page_id are required"}), 400
    record = load_page_extraction(case_id, page_id) or {}
    if not record:
        page = _page_row(load_case_data(case_id), page_id) or {}
        record = {"final_text": str(page.get("page_text", "") or ""), "route": "",
                  "note": "This page was processed before extraction records were kept."}
    return _ok({
        "page_id": page_id,
        "final_text": record.get("final_text", ""),
        "route": record.get("route", ""),
        "text_layer_quality": record.get("text_layer_quality", ""),
        "regions": [{"region": r.get("region"), "kind": r.get("kind"), "flags": r.get("flags") or [],
                     "views": r.get("variants") or [], "numbers_pass": bool(r.get("numbers_pass"))}
                    for r in record.get("regions") or []],
        "candidates": sorted((record.get("candidates") or {}).keys()),
        "uncertain": record.get("uncertain") or [],
        "corrections": record.get("corrections") or [],
        "errors": record.get("errors") or [],
        "document_type": record.get("document_type", ""),
        "note": record.get("note", ""),
    })


def _fact_change_response(case_id, actor):
    result = _after_review_change(case_id, actor)
    data = load_case_data(case_id)
    result["facts_table"] = _fact_table(case_id, data)
    return result


@app.route("/accounting/fact/update", methods=["POST"])
def accounting_fact_update():
    """The user edits cells of one fact (value, description, date, currency,
    type...). Stored as a correction layered on the extraction, merged with
    the user's earlier edits of the same fact; the AI's reading is kept."""
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    row_id = str(body.get("row_id", ""))
    actor = str(body.get("decided_by", "") or "attorney")
    if not (case_id and row_id):
        return jsonify({"error": "case_id and row_id are required."}), 400
    fields, errors = validate_fact_fields(body.get("fields") or {})
    if errors:
        return jsonify({"error": "; ".join(errors)}), 400
    if not fields:
        return jsonify({"error": "Nothing to change."}), 400
    row = _fact_row(case_id, row_id)
    if row is None:
        return jsonify({"error": "That fact no longer exists. Refresh the case."}), 404
    current = effective_fact(row, load_latest_corrections(case_id), load_fact_resolutions(case_id))
    if current["status"] == "IGNORED":
        return jsonify({"error": "Restore the deleted fact before editing it."}), 409
    merged = dict(current.get("final_fields") or {}) if current.get("final_action") == "correct" else {}
    merged.update(fields)
    submit_fact_resolution(case_id, row_id, "correct", merged, str(body.get("note", "") or ""), actor,
                           fingerprint=fact_from_row(row).get("fingerprint", ""))
    return _ok(_fact_change_response(case_id, actor))


@app.route("/accounting/fact/add", methods=["POST"])
def accounting_fact_add():
    """A fact the extraction missed, added by the user on a page."""
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    page_id = str(body.get("page_id", ""))
    actor = str(body.get("decided_by", "") or "attorney")
    if not (case_id and page_id):
        return jsonify({"error": "case_id and page_id are required."}), 400
    fields, errors = validate_fact_fields(body.get("fields") or {})
    if errors:
        return jsonify({"error": "; ".join(errors)}), 400
    if not any(value for name, value in fields.items() if name != "fact_type"):
        return jsonify({"error": "Enter at least a value or a description."}), 400
    data = load_case_data(case_id)
    page = _page_row(data, page_id)
    if not page:
        return jsonify({"error": "That page was not found."}), 404
    document_id = str(page.get("case_document_id", ""))
    names = {str(r.get("case_document_id", "")): r.get("original_filename", "") for r in frame_to_records(data["documents"])}
    row = user_fact_row(page_id, document_id, page.get("page_number", ""), {"fact_type": "other", **fields},
                        names.get(document_id, ""), actor, created_at=pd.Timestamp.utcnow().isoformat())
    persist_reconciled_rows(case_id, [row])
    result = _fact_change_response(case_id, actor)
    result["row_id"] = row["row_id"]
    return _ok(result)


@app.route("/accounting/fact/delete", methods=["POST"])
def accounting_fact_delete():
    """Leave a fact out of the ledger (kept on record, can be restored)."""
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    row_id = str(body.get("row_id", ""))
    actor = str(body.get("decided_by", "") or "attorney")
    row = _fact_row(case_id, row_id) if case_id and row_id else None
    if row is None:
        return jsonify({"error": "That fact no longer exists. Refresh the case."}), 404
    current = effective_fact(row, load_latest_corrections(case_id), load_fact_resolutions(case_id))
    # The user's earlier edits ride along, so a restore brings them back.
    kept = current.get("final_fields") if current.get("final_action") == "correct" else {}
    submit_fact_resolution(case_id, row_id, "ignore", kept or {}, "deleted in the review table", actor,
                           fingerprint=fact_from_row(row).get("fingerprint", ""))
    return _ok(_fact_change_response(case_id, actor))


@app.route("/accounting/fact/restore", methods=["POST"])
def accounting_fact_restore():
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    row_id = str(body.get("row_id", ""))
    actor = str(body.get("decided_by", "") or "attorney")
    row = _fact_row(case_id, row_id) if case_id and row_id else None
    if row is None:
        return jsonify({"error": "That fact no longer exists. Refresh the case."}), 404
    current = effective_fact(row, load_latest_corrections(case_id), load_fact_resolutions(case_id))
    fields = current.get("final_fields") or {}
    fingerprint = fact_from_row(row).get("fingerprint", "")
    if fields:
        submit_fact_resolution(case_id, row_id, "correct", fields, "restored", actor, fingerprint=fingerprint)
    else:
        submit_fact_resolution(case_id, row_id, "confirm", {}, "restored", actor, fingerprint=fingerprint)
    return _ok(_fact_change_response(case_id, actor))


@app.route("/accounting/facts/confirm_all", methods=["POST"])
def accounting_facts_confirm_all():
    """The user has reviewed the fact table: every fact still marked
    uncertain is confirmed as shown, and the accounting analysis unlocks."""
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    actor = str(body.get("decided_by", "") or "attorney")
    if not case_id:
        return jsonify({"error": "case_id is required."}), 400
    frame = case_rows(FINANCIAL_LINE_ITEMS_DATASET, case_id)
    if frame.empty:
        return jsonify({"error": "There are no financial facts to confirm. Run the extraction first."}), 409
    corrections, resolutions = load_latest_corrections(case_id), load_fact_resolutions(case_id)
    decisions = []
    for row in frame.to_dict(orient="records"):
        if effective_fact(row, corrections, resolutions)["status"] == UNCERTAIN:
            decisions.append({"row_id": str(row.get("row_id", "")), "action": "confirm",
                              "explanation": "confirmed with all pages",
                              "fingerprint": fact_from_row(row).get("fingerprint", "")})
    if decisions:
        submit_fact_resolutions(case_id, decisions, actor)
    update_workflow_state(
        case_id,
        changes={"accounting_facts_confirmed": True, "accounting_facts_confirmed_by": actor,
                 "accounting_facts_confirmed_at": pd.Timestamp.utcnow().isoformat()},
        action="accounting_facts_confirmed",
        actor=actor,
    )
    result = _fact_change_response(case_id, actor)
    result["confirmed_uncertain"] = len(decisions)
    return _ok(result)


@app.route("/accounting/fact/explain", methods=["POST"])
def accounting_fact_explain():
    """The user explains how an uncertain fact should be read ("this is the
    third installment"). The AI rewrites the fact from that explanation and
    the page evidence; the rewrite is stored as a proposal that the user
    must accept before it enters the ledger."""
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    row_id = str(body.get("row_id", ""))
    explanation = str(body.get("explanation", "") or "").strip()
    actor = str(body.get("decided_by", "") or "attorney")
    if not (case_id and row_id and explanation):
        return jsonify({"error": "case_id, row_id and an explanation are required."}), 400

    data = load_case_data(case_id)
    row = _fact_row(case_id, row_id)
    if row is None:
        return jsonify({"error": "That fact no longer exists. Refresh the case."}), 404
    current = effective_fact(row, load_latest_corrections(case_id), load_fact_resolutions(case_id))
    page = _pages_by_id(data).get(current["page_id"]) or {}

    try:
        result = reinterpret_fact(current["fact"], current["review"], current["source_text"],
                                  page.get("page_text", ""), explanation)
    except Exception as error:
        traceback.print_exc()
        return jsonify({"error": f"The explanation could not be applied: {error}"}), 502

    fields, errors = validate_fact_fields(result.get("fields") or {})
    if errors:
        return jsonify({"error": "The rewritten fact is not valid: " + "; ".join(errors)}), 502
    summary = str(result.get("summary", "") or "")
    submit_fact_resolution(case_id, row_id, PROPOSAL_ACTION, fields, explanation, actor, summary)
    return _ok({"proposal": {"fields": fields, "summary": summary, "explanation": explanation}})


def _financial_claims(state, data):
    """(claims, source). The attorney review's allegations when a review
    exists; otherwise the alleged facts from the case register, so the
    accounting track does not depend on the attorney review."""
    summary = state.get("attorney_summary") or {}
    allegations = [a for a in (summary.get("allegations") or []) if a]
    if allegations:
        return allegations, "attorney_review"

    claims = [
        {
            "allegation_text": fact["fact_text"],
            "made_by": fact.get("party", ""),
            "source_page_ids": fact.get("source_page_ids", []),
        }
        for fact in case_register(data)["allegations"]
    ]
    return claims, "case_register"


def _claims_to_verify(state, data):
    """(claims, source): the claimant's claims plus the financial claims the
    accounting pages make (e.g. an email saying an amount was transferred).
    Claims are what the accountant verifies; they never enter the ledger."""
    claims, source = _financial_claims(state, data)
    case_id = str((data.get("case") or {}).get("case_id", ""))
    extra = [
        {
            "allegation_text": claim["statement"],
            "made_by": claim.get("made_by") or "document statement",
            "source_page_ids": [claim.get("page_id")] if claim.get("page_id") else [],
            "claim_kind": "financial_claim_in_document",
        }
        for claim in _document_claims(case_id, data)
    ]
    return claims + extra, (source + "+document_claims" if extra else source)


@app.route("/accounting/synthesize", methods=["POST"])
def accounting_synthesize():
    body = request.get_json(force=True)
    case_id = str(body.get("case_id", ""))
    instructions = str(body.get("instructions", "") or "")
    session_id = body.get("session_id")
    if not case_id:
        return jsonify({"error": "case_id is required"}), 400

    data = load_case_data(case_id)
    pending = _review_items(case_id, data)
    if pending:
        return jsonify({
            "error": f"{len(pending)} accounting item(s) are waiting for your confirmation. "
                     "Resolve them before running the analysis.",
            "pending_count": len(pending),
        }), 409
    if not read_workflow_state(case_id).get("accounting_facts_confirmed"):
        return jsonify({"error": "Review the financial facts and press \"Confirm all pages\" before running "
                                 "the accounting analysis."}), 409

    def task(job_id):
        data = load_case_data(case_id)
        state = restore_workflow_state(data)
        normalized_ledger, withheld = _normalized_ledger(case_id, data["financial_line_items"])
        if withheld:
            raise ValueError(f"{withheld} ledger row(s) still await review.")
        if not normalized_ledger:
            raise ValueError("The reviewed ledger is empty. Run the accounting extraction first.")

        claims, claims_source = _claims_to_verify(state, data)

        _set_phase(job_id, "building_timeline", "Building the chronological financial timeline…")
        build_and_save_financial_timeline(case_id, normalized_ledger)

        _set_phase(job_id, "analysing_claims", f"Comparing {len(claims)} claim(s) with the reviewed ledger…")
        result = run_claim_based_accounting_analysis(
            case_id, normalized_ledger, claims, instructions, claims_source=claims_source,
            progress=lambda detail, current, total: _set_phase(
                job_id, "analysing_claims", detail, current=current, total=total),
        )

        if session_id:
            increment_usage(session_id, "llm_request_count", 1)

        update_workflow_state(
            case_id,
            changes={
                "accounting_status": "forensic_complete",
                "accounting_dirty": False,
                "accounting_dirty_reason": "",
                "accounting_instructions": instructions,
            },
            stages={"accounting": stage_patch(
                COMPLETED,
                detail="forensic_complete",
                analysis={"claims": len(claims), "claims_source": claims_source, "ledger_rows": len(normalized_ledger)},
            )},
            action="accounting_synthesized",
        )
        return {"findings": {"claim_evaluations": result.get("claim_evaluations", [])}}

    return _start_stage_job(
        case_id, "accounting", "analysing_claims", "Preparing the accounting analysis…", task, session_id=session_id,
    )


# =============================================================================
# SUMMARY, ANALYSIS & WRITTEN PLEADING ENDPOINTS
# =============================================================================
def _gate_error():
    return (
        "Prepare the consolidated attorney review first."
        if not REQUIRE_APPROVED_SUMMARY
        else "Approve the consolidated attorney review first."
    )


@app.route("/summary/prepare", methods=["POST"])
def summary_prepare():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    session_id = body.get("session_id")

    def task(job_id):
        data = load_case_data(case_id)
        _set_phase(job_id, "preparing_summary", "Preparing consolidated attorney review…")

        # Reads the case register and the documents stage's own page
        # extraction only (pages that recorded claims first), never the
        # accounting classification: the legal track must not depend on
        # whether accounting has run.
        register = case_register(data)
        summary = generate_case_summary(
            case_record=data["case"],
            pages=prioritise_pages_for_summary(data["pages"]),
            facts=register["facts"],
            evidence=data["evidence"],
            parties=data["parties"],
            events=data["events"],
            force_rerun=True,
        )

        if session_id:
            increment_usage(session_id, "llm_request_count", 1)
            increment_usage(session_id, "message_count", 1)

        update_workflow_state(
            case_id,
            changes={"attorney_summary": summary, "summary_approved": False, "summary_approved_by": ""},
            stages={"review": stage_patch(WAITING if REQUIRE_APPROVED_SUMMARY else COMPLETED,
                                          detail="approve_summary" if REQUIRE_APPROVED_SUMMARY else "",
                                          register_fingerprint=register["fingerprint"])},
            action="summary_prepared",
        )
        response_summary = dict(summary)
        if data.get("forensic_findings"):
            response_summary["forensic_findings"] = data["forensic_findings"]

        return {"attorney_summary": response_summary, "summary_support": serialise_summary_support(summary, data)}

    return _start_stage_job(
        case_id, "review", "preparing_summary", "Preparing consolidated attorney review…", task, session_id=session_id,
    )


@app.route("/summary/approve", methods=["POST"])
def summary_approve():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    reviewer = str(body.get("reviewer", "")).strip()
    comments = body.get("comments", "")
    data, state = _load_state(case_id)
    summary = state.get("attorney_summary")
    if not summary or not reviewer:
        return jsonify({"error": "A prepared summary and reviewer name are required."}), 400
    approval_record = approve_case_summary(case_id, summary, reviewer, comments)
    state = update_workflow_state(
        case_id,
        changes={
            "summary_approved": True,
            "summary_approved_by": reviewer,
            "case_dirty": False,
            "dirty_reason": "",
        },
        stages={"review": stage_patch(COMPLETED, detail="approved")},
        action="summary_approved",
        actor=reviewer,
    )
    return _ok({"approval_id": approval_record["approval_id"], "workflow_state": state})


@app.route("/analysis/run", methods=["POST"])
def analysis_run():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    session_id = body.get("session_id")

    def task(job_id):
        data = load_case_data(case_id)
        state = restore_workflow_state(data)
        if not approval_gate_passed(state):
            raise ValueError(_gate_error())
        register = case_register(data)
        facts = register["facts"]
        issues = register["issues"]
        if not facts or not issues:
            raise ValueError("The case register has no facts or issues yet. Process the case documents first.")
        _set_phase(job_id, "researching_law", "Retrieving SAMA authorities and evaluating legal defenses…")
        result = run_legal_analysis(data["case"], facts, issues, register["evidence_requests"])
        if session_id:
            increment_usage(session_id, "llm_request_count", 1)
            increment_usage(session_id, "message_count", 1)
        update_workflow_state(
            case_id,
            changes={"research": result["research"], "analysis": result["analysis"]},
            stages={"analysis": stage_patch(COMPLETED, detail="analysis",
                                            register_fingerprint=register["fingerprint"])},
            action="analysis_prepared",
        )
        return {"research": result["research"], "analysis": result["analysis"]}

    return _start_stage_job(
        case_id, "analysis", "researching_law", "Starting legal analysis…", task, session_id=session_id,
    )


def _drafting_record(case_id, data, state, instructions=""):
    """Everything the written pleading may rely on (see
    legal_platform.pleading.build_drafting_record)."""
    ledger, _ = _normalized_ledger(case_id, data.get("financial_line_items"))
    return build_drafting_record(
        data["case"],
        state.get("attorney_summary"),
        state.get("analysis"),
        state.get("research"),
        data,
        case_register(data),
        _page_reference_map(data),
        ledger,
        ledger_summary(ledger),
        instructions,
    )


@app.route("/pleading/generate", methods=["POST"])
def pleading_generate():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    memo_instructions = body.get("instructions", "")
    direct = bool(body.get("direct", False))
    session_id = body.get("session_id")
    requested_language = str(body.get("language", "") or "").strip().lower()

    def task(job_id):
        data = load_case_data(case_id)
        state = restore_workflow_state(data)
        # The language chosen on the pleading tab (Arabic or English, drafted
        # directly in that language), else the case's preferred one.
        language = requested_language if requested_language in ("ar", "en") else (
            state.get("pleading_language") or str(data["case"].get("preferred_language", "") or "").strip().lower()
            or "en")
        if language not in ("ar", "en"):
            language = "en"

        # 1. Attorney review must be prepared / approved.
        if not approval_gate_passed(state):
            raise ValueError(_gate_error())

        # 2. Legal analysis must be completed.
        if not state.get("analysis"):
            raise ValueError("Complete the legal analysis before generating the written pleading.")

        # 3. Accounting analysis must be completed on a fully reviewed ledger.
        accounting = accounting_state(data, state)
        if accounting["pending_count"]:
            raise ValueError(
                f"{accounting['pending_count']} accounting item(s) are waiting for your confirmation."
            )
        if accounting["status"] not in {"forensic_complete", "complete_no_transactions"}:
            raise ValueError("Complete the accounting analysis before generating the written pleading.")

        _set_phase(job_id, "generating_pleading", "Collecting the record for the pleading…")
        instructions = "" if direct else memo_instructions
        record = _drafting_record(case_id, data, state, instructions)
        memo = draft_pleading(record, language, progress=lambda message: _set_phase(job_id, "generating_pleading", message))
        if session_id:
            increment_usage(session_id, "llm_request_count", 2)
            increment_usage(session_id, "message_count", 1)
        note = "Initial approved-summary draft" if direct else (memo_instructions or "Initial pleading draft")
        versions = [{"version": 1, "draft": memo, "note": note}]
        update_workflow_state(
            case_id,
            changes={
                "memo": memo,
                "pleading_versions": versions,
                "pleading_status": "draft",
                "pleading_finalised_by": "",
                "pleading_language": language,
            },
            stages={"pleading": stage_patch(COMPLETED, detail="draft")},
            action="pleading_prepared",
        )
        return {"memo": memo, "pleading_versions": versions, "pleading_status": "draft"}

    return _start_stage_job(
        case_id, "pleading", "generating_pleading", "Preparing the written pleading…", task, session_id=session_id,
    )


@app.route("/pleading/revise", methods=["POST"])
def pleading_revise():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    revision_request = body.get("revision_request", "")
    session_id = body.get("session_id")

    def task(job_id):
        data = load_case_data(case_id)
        state = restore_workflow_state(data)
        if not state.get("memo"):
            raise ValueError("Generate the written pleading before revising it.")
        _set_phase(job_id, "revising_pleading", "Revising the written pleading…")
        if is_full_pleading(state["memo"]):
            revised = revise_pleading(
                state["memo"], revision_request, _drafting_record(case_id, data, state),
                progress=lambda message: _set_phase(job_id, "revising_pleading", message),
            )
        else:
            # Pleadings drafted before the full structure keep their format.
            revised = revise_bilingual_pleading(
                revision_request,
                state["memo"],
                data["case"],
                state["attorney_summary"],
                state["analysis"],
                None,
                state["research"],
                case_data=data,
            )
        if session_id:
            increment_usage(session_id, "llm_request_count", 1)
            increment_usage(session_id, "message_count", 1)
        versions = list(read_workflow_state(case_id).get("pleading_versions") or [])
        version_no = len(versions) + 1
        versions.append({"version": version_no, "draft": revised, "note": revision_request})
        update_workflow_state(
            case_id,
            changes={"memo": revised, "pleading_versions": versions, "pleading_status": "draft"},
            stages={"pleading": stage_patch(COMPLETED, detail="draft")},
            action="pleading_revised",
            reason=revision_request,
        )
        return {
            "memo": revised,
            "pleading_versions": versions,
            "pleading_status": "draft",
            "version": version_no,
        }

    return _start_stage_job(
        case_id, "pleading", "revising_pleading", "Revising the written pleading…", task, session_id=session_id,
    )


@app.route("/pleading/restore_version", methods=["POST"])
def pleading_restore_version():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    selected_version = body.get("version")
    state = read_workflow_state(case_id)
    versions = list(state.get("pleading_versions") or [])
    selected = next((v for v in versions if v["version"] == selected_version), None)
    if not selected:
        return jsonify({"error": "version not found"}), 404
    version_no = len(versions) + 1
    versions.append({
        "version": version_no, "draft": selected["draft"],
        "note": f"Restored from version {selected_version}",
    })
    state = update_workflow_state(
        case_id,
        changes={"memo": selected["draft"], "pleading_versions": versions},
        action="pleading_version_restored",
    )
    return _ok({"memo": state["memo"], "pleading_versions": state["pleading_versions"]})


@app.route("/pleading/mark_final", methods=["POST"])
def pleading_mark_final():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    final_reviewer = str(body.get("final_reviewer", "")).strip()
    if not final_reviewer:
        return jsonify({"error": "final_reviewer is required"}), 400
    update_workflow_state(
        case_id,
        changes={"pleading_status": "final", "pleading_finalised_by": final_reviewer},
        action="pleading_finalised",
        actor=final_reviewer,
    )
    return jsonify({"pleading_status": "final", "pleading_finalised_by": final_reviewer})


@app.route("/pleading/reopen", methods=["POST"])
def pleading_reopen():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    update_workflow_state(
        case_id,
        changes={"pleading_status": "draft", "pleading_finalised_by": ""},
        action="pleading_reopened",
    )
    return jsonify({"pleading_status": "draft"})


@app.route("/pleading/export")
def pleading_export():
    case_id = request.args.get("case_id", "")
    fmt = request.args.get("fmt", "md")
    language = request.args.get("lang", "en")
    # Only the stored state is needed; the case datasets are not read.
    memo = read_workflow_state(case_id).get("memo")
    if not memo:
        return jsonify({"error": "no pleading available"}), 404
    reference = short_case_reference(case_id)
    if fmt == "docx":
        languages = pleading_languages(memo) if is_full_pleading(memo) else (
            ("ar", "en") if language == "both" else (language,))
        try:
            docx_bytes = build_pleading_docx_bytes(memo, reference, languages)
        except ImportError:
            return jsonify({"error": "python-docx is not installed in this code environment."}), 501
        return send_file(
            io.BytesIO(docx_bytes),
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            as_attachment=True,
            download_name=f"pleading_{reference}_{'_'.join(languages)}.docx",
        )
    text = memo_to_markdown(memo, language)
    return Response(
        text,
        mimetype="text/markdown",
        headers={"Content-Disposition": f'attachment; filename="pleading_{language}.md"'},
    )


@app.route("/discussion/ask", methods=["POST"])
def discussion_ask():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    conversation_id = body.get("conversation_id", "")
    question = body.get("question", "")
    add_to_record = bool(body.get("add_to_record", False))
    session_id = body.get("session_id")

    def task(job_id):
        _set_phase(job_id, "answering", "Answering from the stored case record…")
        data = load_case_data(case_id)
        state = restore_workflow_state(data)
        message_id = add_message(case_id, conversation_id, "user", question)
        memo = state.get("memo")
        answer = answer_case_question(
            question, data["case"], state["attorney_summary"], state["analysis"],
            None, state["research"], case_data=data,
            pleading=memo_to_markdown(memo, pleading_language(memo)) if memo else None,
        )
        increment_usage(session_id, "llm_request_count", 2)
        increment_usage(session_id, "message_count", 1)

        assistant_text = json.dumps({
            "answer_ar": answer.get("answer_ar", ""),
            "answer_en": answer.get("answer_en", ""),
            "uncertainties": answer.get("uncertainties") or [],
            "source_ids": answer.get("source_ids") or [],
        }, ensure_ascii=False)
        add_message(case_id, conversation_id, "assistant", assistant_text)

        if add_to_record:
            add_fact_candidate(
                case_id=case_id, fact_text=question, source_type="chat",
                source_id=message_id, fact_type="user_statement", quote=question, confidence=0.65,
            )
        return {"answer": answer, "invalidated": False}

    return _start_stage_job(
        case_id, "discussion", "answering", "Answering from the stored case record…", task, session_id=session_id,
    )


@app.route("/page_image")
def page_image():
    """One page image. Reads only the page table (not the whole case): the
    review screen requests many images at once. Images never change once
    rendered, so the browser may cache them."""
    case_id = request.args.get("case_id", "")
    page_id = request.args.get("page_id", "")
    row = _page_row({"pages": case_rows("case_document_pages", case_id)}, page_id)
    if not row:
        return Response(status=404)
    image_bytes = _load_page_image_bytes(row.get("page_image_path", ""))
    if not image_bytes:
        return Response(status=404)
    response = send_file(io.BytesIO(image_bytes), mimetype="image/png")
    response.headers["Cache-Control"] = "private, max-age=86400"
    return response


@app.route("/page_text")
def page_text():
    case_id = request.args.get("case_id", "")
    page_id = request.args.get("page_id", "")
    data = load_case_data(case_id)
    row = _page_row(data, page_id)
    if not row:
        return jsonify({"label": "", "text": ""})
    label = _page_reference_map(data).get(str(page_id), "Original page")
    return _ok({"label": label, "text": str(row.get("page_text", "") or "")[:4000]})


@app.route("/flag", methods=["POST"])
def flag_control():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    entity_type = body.get("entity_type", "")
    entity_id = body.get("entity_id", "")
    action = body.get("action", "")
    actor = body.get("actor", "") or "attorney"
    reason = body.get("reason", "")
    if action not in {"verified", "flagged", "corrected"}:
        return jsonify({"error": "invalid action"}), 400
    audit(case_id, entity_type, entity_id, action, actor=actor, reason=reason)
    return jsonify({"ok": True})


# =============================================================================
# AGREEMENT WORKFLOW ENDPOINTS
# =============================================================================
@app.route("/agreement/process", methods=["POST"])
def agreement_process():
    case_id = request.form.get("case_id", "")
    session_id = request.form.get("session_id")
    try:
        zoom = float(request.form.get("zoom", "1.6"))
    except ValueError:
        zoom = 1.6
    uploaded = request.files.getlist("files")
    payloads = [(f.filename, f.read(), f.mimetype) for f in uploaded]
    job_id = _new_job()

    def task():
        processed = 0
        llm_call_count = 0
        for name, raw, content_type in payloads:
            document_row, pdf_bytes = save_upload(case_id, _MemoryUpload(name, raw, content_type), document_type="agreement")
            def on_page(current, total, page_number, state, _name=name):
                _set_progress(job_id, stage="extract", current=current, total=total,
                              detail=f"{_name}: page {page_number}/{total} — {state}")
            pages = extract_pdf_page_by_page(
                case_id=case_id,
                case_document_id=document_row["case_document_id"],
                pdf_bytes=pdf_bytes, zoom=zoom,
                progress_callback=on_page,
            )
            llm_call_count += len(pages)
            if pages:
                processed += 1

        if llm_call_count:
            increment_usage(session_id, "llm_request_count", llm_call_count, use_case="agreement")
        if processed:
            increment_usage(session_id, "attachment_count", processed, use_case="agreement")
        return {"processed": processed}

    _run_async(job_id, task, session_id=session_id, use_case="agreement")
    return jsonify({"job_id": job_id})


@app.route("/agreement/classify", methods=["POST"])
def agreement_classify():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    session_id = body.get("session_id")
    job_id = _new_job()

    def task():
        data = load_case_data(case_id)
        _set_progress(job_id, stage="classify", detail="Detecting agreement type and relationship…")
        result = classify_agreement(data["pages"])
        increment_usage(session_id, "llm_request_count", 1, use_case="agreement")
        increment_usage(session_id, "message_count", 1, use_case="agreement")
        save_agreement_state(case_id, "classification", result)
        return {"classification": result}

    _run_async(job_id, task, session_id=session_id, use_case="agreement")
    return jsonify({"job_id": job_id})


@app.route("/agreement/confirm_classification", methods=["POST"])
def agreement_confirm_classification():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    profile = {
        "agreement_type": body.get("agreement_type"),
        "relationship_type": body.get("relationship_type"),
        "represented_party": str(body.get("represented_party", "")).strip(),
        "counterparty": str(body.get("counterparty", "")).strip(),
        "review_objective": str(body.get("review_objective", "")).strip(),
        "confirmed_by_attorney": True,
    }
    if not profile["represented_party"]:
        return jsonify({"error": "represented_party is required"}), 400
    save_agreement_state(case_id, "confirmed_profile", profile)
    return _ok({"profile": profile})


@app.route("/agreement/extract_clauses", methods=["POST"])
def agreement_extract_clauses():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    session_id = body.get("session_id")
    job_id = _new_job()

    def task():
        data = load_case_data(case_id)
        stored = load_agreement_states(case_id)
        profile = stored.get("confirmed_profile")

        def on_page(current, total, page_id):
            _set_progress(job_id, stage="pages", current=current, total=total,
                          detail=f"Page structure {current}/{total} · {page_id}")

        def on_chunk(current, total, chunk_id):
            _set_progress(job_id, stage="chunks", current=current, total=total,
                          detail=f"Consolidation chunk {current}/{total} · {chunk_id}")

        clause_map = extract_clause_map(
            data["pages"], profile,
            page_progress_callback=on_page,
            chunk_progress_callback=on_chunk,
        )
        increment_usage(session_id, "llm_request_count", 1, use_case="agreement")
        increment_usage(session_id, "message_count", 1, use_case="agreement")
        save_agreement_state(case_id, "clause_map", clause_map)
        return {"clause_map": clause_map}

    _run_async(job_id, task, session_id=session_id, use_case="agreement")
    return jsonify({"job_id": job_id})


@app.route("/agreement/run_review", methods=["POST"])
def agreement_run_review():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    instructions = body.get("instructions", "")
    session_id = body.get("session_id")
    job_id = _new_job()

    def task():
        stored = load_agreement_states(case_id)
        clause_map = stored.get("clause_map")
        profile = stored.get("confirmed_profile")

        def on_review(current, total, clause_id):
            _set_progress(job_id, stage="review", current=current, total=total,
                          detail=f"Clause review {current}/{total} · {clause_id}")

        review, authorities = run_agreement_review(
            clause_map, profile, instructions=instructions, progress_callback=on_review,
        )
        increment_usage(session_id, "llm_request_count", 1, use_case="agreement")
        increment_usage(session_id, "message_count", 1, use_case="agreement")
        authorities_payload = {"authority_nodes": authorities}
        save_agreement_state(case_id, "authorities", authorities_payload)
        save_agreement_state(case_id, "review", review)
        return {"review": review, "authorities": authorities_payload}

    _run_async(job_id, task, session_id=session_id, use_case="agreement")
    return jsonify({"job_id": job_id})


@app.route("/agreement/discuss", methods=["POST"])
def agreement_discuss():
    body = request.get_json(force=True)
    case_id = body.get("case_id", "")
    question = body.get("question", "")
    session_id = body.get("session_id")
    job_id = _new_job()

    def task():
        stored = load_agreement_states(case_id)
        authority_nodes = (stored.get("authorities") or {}).get("authority_nodes", [])
        answer = discuss_agreement(
            question,
            stored.get("confirmed_profile") or {},
            stored.get("clause_map") or {},
            stored.get("review") or {},
            authority_nodes,
        )
        increment_usage(session_id, "llm_request_count", 1, use_case="agreement")
        increment_usage(session_id, "message_count", 1, use_case="agreement")
        return {"answer": answer}

    _run_async(job_id, task, session_id=session_id, use_case="agreement")
    return jsonify({"job_id": job_id})
